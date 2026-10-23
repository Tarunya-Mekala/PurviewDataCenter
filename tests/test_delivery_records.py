"""Tests for Delivery & Installation Records (DocumentType.DELIVERY).

Covers all 4 requirements:
1. DocumentType.DELIVERY support on upload, PATCH, and responses. Refused when source = plan (422).
2. Read files as 1 line = 1 physical device (serial number, PSU serials, hostname, rack, U position, delivery date, DoA test results PASS/FAIL, FAIL sets condition = failed).
3. Used as Found values in comparison against Plan BOM: quantity, model, rack + U position, serial present.
4. Auto-detect files with Serial No. + Rack + RU columns, or DoA/Received/Packing list/Delivery note in sheet or column names.
"""
import io
from pathlib import Path
import pytest
from fastapi import status
import openpyxl


def test_delivery_document_type_upload_refuse_plan(client):
    """Requirement 1: Refuse delivery document type when source = plan."""
    # Create project
    proj_res = client.post("/api/v1/projects", json={"name": "Delivery Plan Refusal Test"})
    assert proj_res.status_code == status.HTTP_201_CREATED
    project_id = proj_res.json()["id"]

    csv_bytes = b"Serial No,Rack,RU,Model\nSN1001,RACK-01,U10,Switch-X\n"
    files = [("files", ("delivery_note_01.csv", io.BytesIO(csv_bytes), "text/csv"))]

    # 1. Explicit document_type='delivery' with source='plan' -> 422
    res_explicit = client.post(
        f"/api/v1/projects/{project_id}/documents",
        files=files,
        data={"source": "plan", "document_type": "delivery"},
    )
    assert res_explicit.status_code == status.HTTP_422_UNPROCESSABLE_ENTITY
    assert "delivery" in res_explicit.json()["detail"].lower()
    assert "vendor" in res_explicit.json()["detail"].lower()


def test_delivery_document_type_upload_and_patch_lifecycle(client):
    """Requirement 1: Delivery accepted on vendor side; PATCH validated."""
    proj_res = client.post("/api/v1/projects", json={"name": "Delivery Lifecycle Test"})
    assert proj_res.status_code == status.HTTP_201_CREATED
    project_id = proj_res.json()["id"]

    # Create vendor
    vnd_res = client.post(f"/api/v1/projects/{project_id}/vendors", json={"name": "Delivery Vendor Inc"})
    assert vnd_res.status_code == status.HTTP_201_CREATED
    vendor_id = vnd_res.json()["id"]

    csv_bytes = b"Serial No,Rack,RU,Model\nSN1001,RACK-01,U10,Switch-X\n"
    files = [("files", ("delivery_record.csv", io.BytesIO(csv_bytes), "text/csv"))]

    # 1. Upload as vendor with document_type='delivery'
    res_upload = client.post(
        f"/api/v1/projects/{project_id}/documents",
        files=files,
        data={"source": "vendor", "vendor_id": vendor_id, "document_type": "delivery"},
    )
    assert res_upload.status_code == status.HTTP_201_CREATED
    doc_info = res_upload.json()["successful"][0]
    assert doc_info["document_type"] == "delivery"
    assert doc_info["source"] == "vendor"
    assert doc_info["vendor_id"] == vendor_id
    doc_id = doc_info["id"]

    # 2. GET /documents/{id} returns document_type == 'delivery'
    get_res = client.get(f"/api/v1/documents/{doc_id}")
    assert get_res.status_code == status.HTTP_200_OK
    assert get_res.json()["document_type"] == "delivery"

    # 3. GET /projects/{id}/documents?document_type=delivery
    list_res = client.get(f"/api/v1/projects/{project_id}/documents?document_type=delivery")
    assert list_res.status_code == status.HTTP_200_OK
    assert list_res.json()["total"] == 1
    assert list_res.json()["items"][0]["document_type"] == "delivery"

    # 4. PATCH to plan source when document_type is delivery -> refused 422
    patch_refuse = client.patch(f"/api/v1/documents/{doc_id}", json={"source": "plan"})
    assert patch_refuse.status_code == status.HTTP_422_UNPROCESSABLE_ENTITY

    # 5. Upload normal document as plan, then PATCH to 'delivery' -> refused 422
    plan_files = [("files", ("normal_plan.csv", io.BytesIO(b"Part Number,Qty\nPN1,1\n"), "text/csv"))]
    plan_upload = client.post(
        f"/api/v1/projects/{project_id}/documents",
        files=plan_files,
        data={"source": "plan", "document_type": "bom"},
    )
    assert plan_upload.status_code == status.HTTP_201_CREATED
    plan_doc_id = plan_upload.json()["successful"][0]["id"]

    patch_plan_to_deliv = client.patch(f"/api/v1/documents/{plan_doc_id}", json={"document_type": "delivery"})
    assert patch_plan_to_deliv.status_code == status.HTTP_422_UNPROCESSABLE_ENTITY


def test_delivery_extraction_device_fields_and_doa_condition(client):
    """Requirement 2: Read delivery files as 1 line = 1 physical device with PSU serials, hostname, date, DoA FAIL sets condition=failed."""
    proj_res = client.post("/api/v1/projects", json={"name": "Delivery Extraction Test"})
    project_id = proj_res.json()["id"]

    vnd_res = client.post(f"/api/v1/projects/{project_id}/vendors", json={"name": "Hardware Integrator"})
    vendor_id = vnd_res.json()["id"]

    # Build Excel file with DoA test results and device attributes
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DoA Testing"
    ws.append([
        "RECEIVED", "Hostname", "PID VID", "SERIAL NO.",
        "Rack", "RU", "PSU PID VID", "PSU1", "PSU2",
        "DOA - PSU", "DOA - FANS", "DOA - CHASSIS", "DOA - LEDs", "NOTES"
    ])
    # Device 1: PASS
    ws.append([
        "2026-05-20", "leaf-sw-01", "HF6100-32D V01", "FLM3001001",
        "RACK-A01", "U12", "PSU-3000W", "DTM101A", "DTM101B",
        "PASS", "PASS", "PASS", "PASS", None
    ])
    # Device 2: FAIL on FANS
    ws.append([
        "2026-05-20", "leaf-sw-02", "HF6100-32D V01", "FLM3001002",
        "RACK-A01", "U13", "PSU-3000W", "DTM102A", "DTM102B",
        "PASS", "FAIL", "PASS", "PASS", "FANS NOT SPINNING"
    ])

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    upload_res = client.post(
        f"/api/v1/projects/{project_id}/documents",
        files=[("files", ("delivery_doa.xlsx", bio, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))],
        data={"source": "vendor", "vendor_id": vendor_id, "document_type": "delivery"},
    )
    assert upload_res.status_code == status.HTTP_201_CREATED
    doc_id = upload_res.json()["successful"][0]["id"]

    # Extract records
    extract_res = client.post(f"/api/v1/documents/{doc_id}/extract")
    assert extract_res.status_code == status.HTTP_200_OK
    assert extract_res.json()["records_extracted"] == 2

    # Verify extracted records
    records_res = client.get(f"/api/v1/projects/{project_id}/records")
    items = records_res.json()["items"]
    assert len(items) == 2

    # Map by serial
    by_serial = {r["serial_number"]: r for r in items}

    # Verify Device 1 (Healthy PASS)
    dev1 = by_serial["FLM3001001"]
    assert dev1["part_number"] == "HF6100-32D V01"
    assert dev1["cleaned_part_number"] == "HF6100-32D"
    assert dev1["model"] == "HF6100-32D"
    assert dev1["rack"] == "RACK-A01"
    assert dev1["u_position"] == "U12"
    assert dev1["hostname"] == "leaf-sw-01"
    assert dev1["delivery_date"] == "2026-05-20"
    assert "DTM101A" in (dev1["psu_serials"] or "")
    assert dev1["condition"] == "ok"

    # Verify Device 2 (FAIL on FANS)
    dev2 = by_serial["FLM3001002"]
    assert dev2["part_number"] == "HF6100-32D V01"
    assert dev2["cleaned_part_number"] == "HF6100-32D"
    assert dev2["rack"] == "RACK-A01"
    assert dev2["u_position"] == "U13"
    assert dev2["hostname"] == "leaf-sw-02"
    assert dev2["condition"] == "failed"
    assert "FANS" in (dev2["condition_notes"] or "")


def test_delivery_used_as_found_values_in_comparison(client):
    """Requirement 3: Use delivery records as Found values in comparison against Plan BOM."""
    proj_res = client.post("/api/v1/projects", json={"name": "Delivery Comparison Test"})
    project_id = proj_res.json()["id"]

    vnd_res = client.post(f"/api/v1/projects/{project_id}/vendors", json={"name": "Sharon AI Systems"})
    vendor_id = vnd_res.json()["id"]

    # 1. Upload Plan BOM: 2 units of HF6100-32D in RACK-A01
    plan_csv = (
        "Part Number,Description,Qty,Rack,U Position\n"
        "HF6100-32D,Storage Leaf Switch,2,RACK-A01,U12\n"
    )
    plan_upload = client.post(
        f"/api/v1/projects/{project_id}/documents",
        files=[("files", ("plan_bom.csv", io.BytesIO(plan_csv.encode()), "text/csv"))],
        data={"source": "plan", "document_type": "bom"},
    )
    plan_doc_id = plan_upload.json()["successful"][0]["id"]
    client.post(f"/api/v1/documents/{plan_doc_id}/extract")

    # 2. Upload Delivery record: 2 physical units arrived
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Delivery Note"
    ws.append(["Part Number", "Description", "SERIAL NO.", "Rack", "RU", "DOA - PSU", "DOA - FANS", "NOTES"])
    ws.append(["HF6100-32D", "Storage Leaf Switch", "SER-001", "RACK-A01", "U12", "PASS", "PASS", None])
    ws.append(["HF6100-32D", "Storage Leaf Switch", "SER-002", "RACK-A01", "U12", "PASS", "PASS", None])

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    deliv_upload = client.post(
        f"/api/v1/projects/{project_id}/documents",
        files=[("files", ("delivery_notes.xlsx", bio, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))],
        data={"source": "vendor", "vendor_id": vendor_id, "document_type": "delivery"},
    )
    deliv_doc_id = deliv_upload.json()["successful"][0]["id"]
    client.post(f"/api/v1/documents/{deliv_doc_id}/extract")

    # Approve all records
    records = client.get(f"/api/v1/projects/{project_id}/records").json()["items"]
    for r in records:
        client.post(f"/api/v1/records/{r['id']}/review", json={"status": "approved", "review_notes": "Approved for comparison"})

    # 3. Run comparison
    cmp_res = client.post(f"/api/v1/projects/{project_id}/compare")
    assert cmp_res.status_code == status.HTTP_200_OK
    cmp_data = cmp_res.json()

    # Verify Expected vs Found
    summary = cmp_data["summary"]
    assert summary["total_expected"] == 2.0
    assert summary["total_found"] == 2.0
    assert summary["matches_count"] == 1

    item = cmp_data["items"][0]
    assert item["expected_quantity"] == 2.0
    assert item["found_quantity"] == 2.0
    assert item["status"] == "MATCH"
    assert "Serial present" in item["details"]
    assert "SER-001" in item["details"]


def test_delivery_auto_detection_from_filename_and_content(client):
    """Requirement 4: Auto-detect files with Serial No. + Rack + RU columns, or DoA/Received/Packing list/Delivery note."""
    proj_res = client.post("/api/v1/projects", json={"name": "Auto Detection Test"})
    project_id = proj_res.json()["id"]

    vnd_res = client.post(f"/api/v1/projects/{project_id}/vendors", json={"name": "Vendor 1"})
    vendor_id = vnd_res.json()["id"]

    # 1. Filename auto-detect: "SharonAI Device Hostname_Location_Serial Numbers.xlsx"
    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.append(["Device", "Hostname", "Rack", "RU", "Serial No."])
    ws1.append(["HF6100", "sw01", "R01", "U10", "SN-999"])
    b1 = io.BytesIO()
    wb1.save(b1)
    b1.seek(0)

    up1 = client.post(
        f"/api/v1/projects/{project_id}/documents",
        files=[("files", ("SharonAI Device Hostname_Location_Serial Numbers.xlsx", b1, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))],
        data={"source": "vendor", "vendor_id": vendor_id},  # No category given!
    )
    assert up1.status_code == status.HTTP_201_CREATED
    assert up1.json()["successful"][0]["document_type"] == "delivery"

    # 2. Filename auto-detect: "SharonAI Active Equipment - DoA Testing.xlsx"
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.append(["Device", "Serial", "DOA Status"])
    ws2.append(["PDU", "SN-888", "PASS"])
    b2 = io.BytesIO()
    wb2.save(b2)
    b2.seek(0)

    up2 = client.post(
        f"/api/v1/projects/{project_id}/documents",
        files=[("files", ("SharonAI Active Equipment - DoA Testing.xlsx", b2, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))],
        data={"source": "vendor", "vendor_id": vendor_id},  # No category given!
    )
    assert up2.status_code == status.HTTP_201_CREATED
    assert up2.json()["successful"][0]["document_type"] == "delivery"

    # 3. Content auto-detect: generic filename "hardware_records.xlsx" with sheet named "Packing list"
    wb3 = openpyxl.Workbook()
    ws3 = wb3.active
    ws3.title = "Packing list"
    ws3.append(["Item", "Serial Number", "Rack", "RU"])
    ws3.append(["Server 1", "SN-777", "R02", "U20"])
    b3 = io.BytesIO()
    wb3.save(b3)
    b3.seek(0)

    up3 = client.post(
        f"/api/v1/projects/{project_id}/documents",
        files=[("files", ("hardware_records.xlsx", b3, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))],
        data={"source": "vendor", "vendor_id": vendor_id},  # No category given!
    )
    assert up3.status_code == status.HTTP_201_CREATED
    assert up3.json()["successful"][0]["document_type"] == "delivery"

    # 4. Content auto-detect: generic CSV with Serial No. + Rack + RU columns
    csv_bytes = b"Serial No.,Cabinet,Rack Unit,Model\nSN-555,CAB-1,U04,Switch-Z\n"
    up4 = client.post(
        f"/api/v1/projects/{project_id}/documents",
        files=[("files", ("unnamed_sheet.csv", io.BytesIO(csv_bytes), "text/csv"))],
        data={"source": "vendor", "vendor_id": vendor_id},  # No category given!
    )
    assert up4.status_code == status.HTTP_201_CREATED
    assert up4.json()["successful"][0]["document_type"] == "delivery"

    # 5. When auto-detected as delivery on source = plan -> refused with 422!
    up5_refuse = client.post(
        f"/api/v1/projects/{project_id}/documents",
        files=[("files", ("SharonAI Active Equipment - DoA Testing.xlsx", io.BytesIO(b2.getvalue()), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))],
        data={"source": "plan"},  # No category given, but source is plan!
    )
    assert up5_refuse.status_code == status.HTTP_422_UNPROCESSABLE_ENTITY
    assert "delivery" in up5_refuse.json()["detail"].lower()
