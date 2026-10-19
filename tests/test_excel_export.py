import io
import uuid
import openpyxl
from fastapi import status


def _setup_extracted_document(client):
    """Helper to create a project, upload a BOM file, and extract records."""
    # 1. Create project with unique name
    proj_name = f"Excel Export Test Project {uuid.uuid4().hex[:8]}"
    proj_res = client.post("/api/v1/projects", json={"name": proj_name})
    assert proj_res.status_code == status.HTTP_201_CREATED
    project_id = proj_res.json()["id"]

    # 2. Upload CSV BOM
    bom_csv = (
        "Part Number,Description,Qty,Vendor,Rack\n"
        "N9K-C93180YC-FX,Nexus 9300 48-port 25G Switch,2,Cisco,Rack-A01\n"
        "DCS-7050TX3-48C8,Arista 48x10G 8x100G Switch,4,Arista,Rack-A02\n"
        "PowerEdge-R750,Dell Dual Xeon Server 64GB,8,Dell,Rack-B01\n"
    ).encode("utf-8")

    files = [("files", ("test_datacenter_bom.csv", io.BytesIO(bom_csv), "text/csv"))]
    upload_res = client.post(f"/api/v1/projects/{project_id}/documents", files=files)
    assert upload_res.status_code == status.HTTP_201_CREATED
    doc_id = upload_res.json()["successful"][0]["id"]

    # 3. Extract document
    extract_res = client.post(f"/api/v1/documents/{doc_id}/extract")
    assert extract_res.status_code == status.HTTP_200_OK
    assert extract_res.json()["records_extracted"] == 3

    return project_id, doc_id


def test_document_excel_download(client):
    project_id, doc_id = _setup_extracted_document(client)

    # 1. Download Excel workbook (.xlsx)
    res = client.get(f"/api/v1/documents/{doc_id}/excel")
    assert res.status_code == status.HTTP_200_OK
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
    assert "attachment" in res.headers["content-disposition"]
    assert ".xlsx" in res.headers["content-disposition"]

    # 2. Validate that openpyxl can parse the generated workbook
    excel_content = res.content
    assert len(excel_content) > 0
    wb = openpyxl.load_workbook(io.BytesIO(excel_content), data_only=False)
    assert "Extracted Records" in wb.sheetnames
    assert "Document Overview" in wb.sheetnames

    # Check Extracted Records sheet content
    ws_records = wb["Extracted Records"]
    assert ws_records.cell(row=1, column=1).value == "Line #"
    assert ws_records.cell(row=1, column=3).value == "Part Number"
    assert ws_records.cell(row=1, column=5).value == "Quantity"
    
    # Check data rows
    assert ws_records.cell(row=2, column=1).value == 1
    assert "N9K" in str(ws_records.cell(row=2, column=3).value)
    assert ws_records.cell(row=2, column=5).value == 2.0
    assert ws_records.cell(row=2, column=9).value == "Cisco"

    # Check Document Overview sheet
    ws_overview = wb["Document Overview"]
    assert "Extraction Overview" in str(ws_overview.cell(row=1, column=1).value)
    wb.close()


def test_document_export_alias_and_reports_endpoint(client):
    project_id, doc_id = _setup_extracted_document(client)

    # Direct export alias
    res1 = client.get(f"/api/v1/documents/{doc_id}/export/excel")
    assert res1.status_code == status.HTTP_200_OK
    assert "spreadsheetml.sheet" in res1.headers["content-type"]

    # Reports router endpoint
    res2 = client.get(f"/api/v1/reports/documents/{doc_id}/excel")
    assert res2.status_code == status.HTTP_200_OK
    assert "spreadsheetml.sheet" in res2.headers["content-type"]


def test_document_excel_html_view(client):
    project_id, doc_id = _setup_extracted_document(client)

    # 1. Access dedicated /excel/view endpoint
    res = client.get(f"/api/v1/documents/{doc_id}/excel/view")
    assert res.status_code == status.HTTP_200_OK
    assert "text/html" in res.headers["content-type"]
    html = res.text
    assert "test_datacenter_bom.csv" in html
    assert "N9K-C93180YC-FX" in html
    assert "Download .xlsx" in html
    assert "Total Line Items" in html

    # 2. Access /excel with view=html query parameter
    res_query = client.get(f"/api/v1/documents/{doc_id}/excel?view=html")
    assert res_query.status_code == status.HTTP_200_OK
    assert "text/html" in res_query.headers["content-type"]

    # 3. Access /reports/documents/{doc_id}/excel/view
    res_report_view = client.get(f"/api/v1/reports/documents/{doc_id}/excel/view")
    assert res_report_view.status_code == status.HTTP_200_OK
    assert "text/html" in res_report_view.headers["content-type"]


def test_document_excel_preview_json(client):
    project_id, doc_id = _setup_extracted_document(client)

    res = client.get(f"/api/v1/documents/{doc_id}/excel/preview")
    assert res.status_code == status.HTTP_200_OK
    data = res.json()
    assert data["document_id"] == doc_id
    assert data["sheet_name"] == "Extracted Records"
    assert data["total_records"] == 3
    assert "columns" in data
    assert "rows" in data
    assert len(data["rows"]) == 3
    assert data["summary"]["total"] == 3
    assert data["summary"]["pending"] == 3


def test_project_excel_download_and_views(client):
    project_id, doc_id = _setup_extracted_document(client)

    # 1. Project Excel download (.xlsx)
    res = client.get(f"/api/v1/projects/{project_id}/excel")
    assert res.status_code == status.HTTP_200_OK
    assert "spreadsheetml.sheet" in res.headers["content-type"]
    assert "attachment" in res.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(res.content), data_only=False)
    assert "All Project Records" in wb.sheetnames
    assert "Documents Summary" in wb.sheetnames
    assert "Project Overview" in wb.sheetnames
    wb.close()

    # 2. Project export alias
    res_alias = client.get(f"/api/v1/projects/{project_id}/export/excel")
    assert res_alias.status_code == status.HTTP_200_OK

    # 3. Project Excel HTML view
    res_view = client.get(f"/api/v1/projects/{project_id}/excel/view")
    assert res_view.status_code == status.HTTP_200_OK
    assert "text/html" in res_view.headers["content-type"]
    assert "Consolidated Project Records" in res_view.text

    # 4. Project Excel JSON preview
    res_prev = client.get(f"/api/v1/projects/{project_id}/excel/preview")
    assert res_prev.status_code == status.HTTP_200_OK
    assert res_prev.json()["total_records"] == 3


def test_token_query_param_auth(unauth_client, auth_token, client):
    project_id, doc_id = _setup_extracted_document(client)

    # 1. Requesting without token should fail 401
    unauth_res = unauth_client.get(f"/api/v1/documents/{doc_id}/excel")
    assert unauth_res.status_code == status.HTTP_401_UNAUTHORIZED

    # 2. Requesting with ?token= query parameter should succeed
    token_res = unauth_client.get(f"/api/v1/documents/{doc_id}/excel?token={auth_token}")
    assert token_res.status_code == status.HTTP_200_OK
    assert "spreadsheetml.sheet" in token_res.headers["content-type"]

    # 3. HTML view with ?token= query parameter
    token_view_res = unauth_client.get(f"/api/v1/documents/{doc_id}/excel/view?token={auth_token}")
    assert token_view_res.status_code == status.HTTP_200_OK
    assert "text/html" in token_view_res.headers["content-type"]


def test_excel_export_not_found(client):
    res_doc = client.get("/api/v1/documents/non_existent_doc/excel")
    assert res_doc.status_code == status.HTTP_404_NOT_FOUND

    res_proj = client.get("/api/v1/projects/non_existent_proj/excel")
    assert res_proj.status_code == status.HTTP_404_NOT_FOUND


def test_empty_document_excel_export(client):
    # Create project and empty document (no records extracted yet)
    proj_res = client.post("/api/v1/projects", json={"name": "Empty Doc Test"})
    project_id = proj_res.json()["id"]

    empty_csv = b"Part Number,Description\n"
    files = [("files", ("empty_bom.csv", io.BytesIO(empty_csv), "text/csv"))]
    upload_res = client.post(f"/api/v1/projects/{project_id}/documents", files=files)
    doc_id = upload_res.json()["successful"][0]["id"]

    # Export document with 0 records
    res = client.get(f"/api/v1/documents/{doc_id}/excel")
    assert res.status_code == status.HTTP_200_OK
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "Extracted Records" in wb.sheetnames
    wb.close()
