import io
import uuid
import openpyxl
from fastapi import status
from fastapi.testclient import TestClient

from app.core.database import get_db
from app.core.security import create_access_token, hash_password
from app.main import app
from app.models.user import User
from app.services.report_service import ReportService


def _setup_test_document(client):
    proj_name = f"Export Protection Proj {uuid.uuid4().hex[:8]}"
    proj_res = client.post("/api/v1/projects", json={"name": proj_name})
    assert proj_res.status_code == status.HTTP_201_CREATED
    project_id = proj_res.json()["id"]

    bom_csv = (
        "Part Number,Description,Qty,Vendor,Rack\n"
        "SW-101,Core Switch 48-port,2,Cisco,Rack-1\n"
        "SRV-201,Compute Node 128G,4,Dell,Rack-2\n"
    ).encode("utf-8")

    files = [("files", ("switch_bom.csv", io.BytesIO(bom_csv), "text/csv"))]
    upload_res = client.post(f"/api/v1/projects/{project_id}/documents", files=files)
    assert upload_res.status_code == status.HTTP_201_CREATED
    doc_id = upload_res.json()["successful"][0]["id"]

    extract_res = client.post(f"/api/v1/documents/{doc_id}/extract")
    assert extract_res.status_code == status.HTTP_200_OK

    return project_id, doc_id


def _create_user(db_session, email_prefix, role, is_super_admin=False):
    user = User(
        email=f"{email_prefix}_{uuid.uuid4().hex[:6]}@example.com",
        hashed_password=hash_password("TestPassword123!"),
        full_name=f"Test {role.title()}",
        role=role,
        is_active=True,
        is_super_admin=is_super_admin,
    )
    db_session.add(user)
    db_session.commit()
    db_session.refresh(user)

    token = create_access_token(
        data={
            "sub": user.id,
            "email": user.email,
            "role": user.role,
            "is_super_admin": user.is_super_admin,
        }
    )
    return user, token


def test_file_and_title_renamed_to_internal_data_export(client):
    project_id, doc_id = _setup_test_document(client)

    # 1. Check document Excel filename and sheet title
    res_doc = client.get(f"/api/v1/reports/documents/{doc_id}/excel")
    assert res_doc.status_code == status.HTTP_200_OK
    content_disp = res_doc.headers.get("content-disposition", "")
    assert "Extracted_lines_-_internal_data_export_" in content_disp

    wb = openpyxl.load_workbook(io.BytesIO(res_doc.content))
    ws_overview = wb["Document Overview"]
    title_val = str(ws_overview.cell(row=1, column=1).value)
    assert "Extracted lines – internal data export" in title_val
    wb.close()

    # 2. Check document HTML view title
    res_html = client.get(f"/api/v1/reports/documents/{doc_id}/excel/view")
    assert res_html.status_code == status.HTTP_200_OK
    assert "<title>Extracted lines – internal data export" in res_html.text
    assert "<h1>Extracted lines – internal data export</h1>" in res_html.text

    # 3. Check project Excel filename and sheet title
    res_proj = client.get(f"/api/v1/reports/projects/{project_id}/excel")
    assert res_proj.status_code == status.HTTP_200_OK
    proj_disp = res_proj.headers.get("content-disposition", "")
    assert "Extracted_lines_-_internal_data_export_" in proj_disp

    wb_proj = openpyxl.load_workbook(io.BytesIO(res_proj.content))
    ws_proj_overview = wb_proj["Project Overview"]
    proj_title_val = str(ws_proj_overview.cell(row=1, column=1).value)
    assert "Extracted lines – internal data export" in proj_title_val
    wb_proj.close()

    # 4. Check project HTML view title
    res_proj_html = client.get(f"/api/v1/reports/projects/{project_id}/excel/view")
    assert res_proj_html.status_code == status.HTTP_200_OK
    assert "<title>Extracted lines – internal data export" in res_proj_html.text
    assert "<h1>Extracted lines – internal data export</h1>" in res_proj_html.text


def test_rbac_reviewer_and_admin_allowed(client, db_session):
    project_id, doc_id = _setup_test_document(client)

    def _client_with_token(token):
        def override_get_db():
            try:
                yield db_session
            finally:
                pass

        app.dependency_overrides[get_db] = override_get_db
        test_client = TestClient(app)
        test_client.headers.update({"Authorization": f"Bearer {token}"})
        return test_client

    # Create reviewer user
    _, reviewer_token = _create_user(db_session, "reviewer", "reviewer", is_super_admin=False)
    reviewer_client = _client_with_token(reviewer_token)

    # Reviewer should have access to /reports/.../excel, /view, and /preview
    assert reviewer_client.get(f"/api/v1/reports/documents/{doc_id}/excel").status_code == status.HTTP_200_OK
    assert reviewer_client.get(f"/api/v1/reports/documents/{doc_id}/excel/view").status_code == status.HTTP_200_OK
    assert reviewer_client.get(f"/api/v1/reports/documents/{doc_id}/excel/preview").status_code == status.HTTP_200_OK
    assert reviewer_client.get(f"/api/v1/reports/projects/{project_id}/excel").status_code == status.HTTP_200_OK
    assert reviewer_client.get(f"/api/v1/reports/projects/{project_id}/excel/view").status_code == status.HTTP_200_OK
    assert reviewer_client.get(f"/api/v1/reports/projects/{project_id}/excel/preview").status_code == status.HTTP_200_OK

    # Create admin user
    _, admin_token = _create_user(db_session, "admin", "admin", is_super_admin=False)
    admin_client = _client_with_token(admin_token)

    # Admin should have access to all endpoints
    assert admin_client.get(f"/api/v1/reports/documents/{doc_id}/excel").status_code == status.HTTP_200_OK
    assert admin_client.get(f"/api/v1/reports/documents/{doc_id}/excel/view").status_code == status.HTTP_200_OK
    assert admin_client.get(f"/api/v1/reports/documents/{doc_id}/excel/preview").status_code == status.HTTP_200_OK
    assert admin_client.get(f"/api/v1/reports/projects/{project_id}/excel").status_code == status.HTTP_200_OK
    assert admin_client.get(f"/api/v1/reports/projects/{project_id}/excel/view").status_code == status.HTTP_200_OK
    assert admin_client.get(f"/api/v1/reports/projects/{project_id}/excel/preview").status_code == status.HTTP_200_OK


def test_rbac_non_reviewer_non_admin_forbidden(client, db_session):
    project_id, doc_id = _setup_test_document(client)

    def _client_with_token(token):
        def override_get_db():
            try:
                yield db_session
            finally:
                pass

        app.dependency_overrides[get_db] = override_get_db
        test_client = TestClient(app)
        test_client.headers.update({"Authorization": f"Bearer {token}"})
        return test_client

    # Create engineer/viewer user
    _, engineer_token = _create_user(db_session, "engineer", "engineer", is_super_admin=False)
    engineer_client = _client_with_token(engineer_token)

    # Non-reviewer/non-admin should get 403 Forbidden
    assert engineer_client.get(f"/api/v1/reports/documents/{doc_id}/excel").status_code == status.HTTP_403_FORBIDDEN
    assert engineer_client.get(f"/api/v1/reports/documents/{doc_id}/excel/view").status_code == status.HTTP_403_FORBIDDEN
    assert engineer_client.get(f"/api/v1/reports/documents/{doc_id}/excel/preview").status_code == status.HTTP_403_FORBIDDEN
    assert engineer_client.get(f"/api/v1/reports/projects/{project_id}/excel").status_code == status.HTTP_403_FORBIDDEN
    assert engineer_client.get(f"/api/v1/reports/projects/{project_id}/excel/view").status_code == status.HTTP_403_FORBIDDEN
    assert engineer_client.get(f"/api/v1/reports/projects/{project_id}/excel/preview").status_code == status.HTTP_403_FORBIDDEN


def test_unauthenticated_rejected(unauth_client, client):
    project_id, doc_id = _setup_test_document(client)

    assert unauth_client.get(f"/api/v1/reports/documents/{doc_id}/excel").status_code == status.HTTP_401_UNAUTHORIZED
    assert unauth_client.get(f"/api/v1/reports/documents/{doc_id}/excel/view").status_code == status.HTTP_401_UNAUTHORIZED
    assert unauth_client.get(f"/api/v1/reports/documents/{doc_id}/excel/preview").status_code == status.HTTP_401_UNAUTHORIZED
    assert unauth_client.get(f"/api/v1/reports/projects/{project_id}/excel").status_code == status.HTTP_401_UNAUTHORIZED
    assert unauth_client.get(f"/api/v1/reports/projects/{project_id}/excel/view").status_code == status.HTTP_401_UNAUTHORIZED
    assert unauth_client.get(f"/api/v1/reports/projects/{project_id}/excel/preview").status_code == status.HTTP_401_UNAUTHORIZED


def test_html_view_signed_download_link_and_omission_without_token(client, unauth_client, db_session):
    project_id, doc_id = _setup_test_document(client)

    # 1. When accessed via authenticated endpoint, HTML view includes signed ?token= download link
    res_html = client.get(f"/api/v1/reports/documents/{doc_id}/excel/view")
    assert res_html.status_code == status.HTTP_200_OK
    assert "Download .xlsx" in res_html.text
    assert f"/api/v1/documents/{doc_id}/excel?token=" in res_html.text

    # Extract the token from the link and verify unauthenticated browser download succeeds with it
    html_text = res_html.text
    token_start = html_text.find(f"/api/v1/documents/{doc_id}/excel?token=") + len(f"/api/v1/documents/{doc_id}/excel?token=")
    token_end = html_text.find('"', token_start)
    extracted_token = html_text[token_start:token_end]
    assert len(extracted_token) > 20

    download_res = unauth_client.get(f"/api/v1/documents/{doc_id}/excel?token={extracted_token}")
    assert download_res.status_code == status.HTTP_200_OK
    assert "spreadsheetml.sheet" in download_res.headers["content-type"]

    # 2. When rendered directly with token=None (e.g. unauthenticated view rendering), the button is omitted
    raw_html_no_token = ReportService.render_excel_html_view(db_session, doc_id, token=None)
    assert "Download .xlsx" not in raw_html_no_token
    assert 'id="download-excel-btn"' not in raw_html_no_token

    # 3. Check project HTML view signed link
    res_proj_html = client.get(f"/api/v1/reports/projects/{project_id}/excel/view")
    assert res_proj_html.status_code == status.HTTP_200_OK
    assert "Download Project .xlsx" in res_proj_html.text
    assert f"/api/v1/projects/{project_id}/excel?token=" in res_proj_html.text

    raw_proj_no_token = ReportService.render_project_excel_html_view(db_session, project_id, token=None)
    assert "Download Project .xlsx" not in raw_proj_no_token
