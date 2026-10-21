from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl


EXCEL_EXTENSIONS = {
    ".xlsx",
    ".xlsm",
    ".xltx",
    ".xltm",
}


def _clean_value(value: Any) -> Any:
    """
    Convert Excel cell values into JSON/text friendly values.
    """

    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()
        return value if value else None

    return value


def _column_letter(column_number: int) -> str:
    """
    Convert 1-based column number to Excel column name.

    1 -> A
    2 -> B
    27 -> AA
    """

    result = ""

    while column_number:
        column_number, remainder = divmod(
            column_number - 1,
            26,
        )

        result = (
            chr(65 + remainder)
            + result
        )

    return result


def parse_excel_file(
    file_path: Path,
) -> dict[str, Any]:
    """
    Parse the COMPLETE Excel workbook.

    No fixed:
        - sheet count
        - row count
        - column count
        - column names
        - column positions

    Every populated cell is preserved.

    Output contains:

        workbook
            sheets
                rows
                    cells
    """

    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"Excel file not found: {file_path}"
        )

    if file_path.suffix.lower() not in EXCEL_EXTENSIONS:
        raise ValueError(
            f"Unsupported Excel extension: "
            f"{file_path.suffix}"
        )

    workbook = openpyxl.load_workbook(
        filename=file_path,
        data_only=True,
        read_only=True,
    )

    sheets: list[dict[str, Any]] = []

    try:

        for sheet in workbook.worksheets:

            rows: list[dict[str, Any]] = []

            for row_number, row in enumerate(
                sheet.iter_rows(),
                start=1,
            ):

                cells: list[dict[str, Any]] = []

                has_value = False

                for column_number, cell in enumerate(
                    row,
                    start=1,
                ):

                    value = _clean_value(
                        cell.value
                    )

                    if value is not None:
                        has_value = True

                    cells.append(
                        {
                            "column_number": column_number,
                            "column_letter": _column_letter(
                                column_number
                            ),
                            "value": value,
                        }
                    )

                # Preserve the row even when some
                # cells are empty.
                if has_value:
                    rows.append(
                        {
                            "row_number": row_number,
                            "cells": cells,
                        }
                    )

            sheets.append(
                {
                    "sheet_name": sheet.title,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                    "rows": rows,
                }
            )

    finally:
        workbook.close()

    return {
        "file_name": file_path.name,
        "file_type": file_path.suffix.lower(),
        "sheet_count": len(sheets),
        "sheets": sheets,
    }


def excel_to_llm_text(
    parsed_workbook: dict[str, Any],
) -> str:
    """
    Convert the structured Excel representation into
    an explicit representation for the LLM.

    The important thing is that column positions and
    row numbers are preserved.
    """

    output: list[str] = []

    output.append(
        "===== EXCEL WORKBOOK ====="
    )

    output.append(
        f"FILE: "
        f"{parsed_workbook.get('file_name', '')}"
    )

    output.append(
        f"SHEETS: "
        f"{parsed_workbook.get('sheet_count', 0)}"
    )

    for sheet in parsed_workbook.get(
        "sheets",
        [],
    ):

        output.append(
            f"\n===== SHEET: "
            f"{sheet['sheet_name']} ====="
        )

        output.append(
            f"MAX_ROWS: {sheet.get('max_row')}"
        )

        output.append(
            f"MAX_COLUMNS: "
            f"{sheet.get('max_column')}"
        )

        for row in sheet.get(
            "rows",
            [],
        ):

            row_number = row[
                "row_number"
            ]

            output.append(
                f"\nROW {row_number}:"
            )

            for cell in row.get(
                "cells",
                [],
            ):

                output.append(
                    f"  COLUMN "
                    f"{cell['column_letter']} "
                    f"(index "
                    f"{cell['column_number']}): "
                    f"{cell['value']}"
                )

    return "\n".join(output)