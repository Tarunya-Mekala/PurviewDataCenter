import csv
import io
import logging
import random
import re
from decimal import Decimal, InvalidOperation
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from fastapi import BackgroundTasks, HTTPException, status

logger = logging.getLogger(__name__)
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.extraction.column_map import (
    clean_part_number,
    detect_condition,
    map_headers,
    normalize_header,
)
from app.models.document import Document
from app.models.extracted_record import ExtractedRecord
from app.models.project import Project
from app.models.record_review_action import RecordReviewAction
from app.schemas.document import DocumentStatus
from app.schemas.record import (
    BatchReviewAction,
    BatchReviewRequest,
    DocumentExtractionStatusResponse,
    ExtractedRecordCreate,
    ExtractedRecordReview,
    ExtractedRecordUpdate,
    ExtractionJobResponse,
    RecordType,
    ReviewStatus,
    ReviewSummaryResponse,
)
from app.utils.id_generator import generate_id
from app.utils.normalization import extract_vendor_from_text


def col_num_to_letter(col_idx: int) -> str:
    """Convert 0-indexed column integer to Excel column letter (0->A, 1->B, 26->AA)."""
    n = col_idx + 1
    result = ""
    while n > 0:
        n, remainder = divmod(n - 1, 26)
        result = chr(65 + remainder) + result
    return result


def parse_quantity_strict(val: Any) -> Optional[float]:
    """Safely convert quantity to float if present; return None if empty/null/missing or date."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if not s:
        return None
    # Never map dates (e.g. 25.05.26, 2026-05-25, 25/05/2026) to quantity
    if re.search(r"^\d{1,4}[./-]\d{1,2}[./-]\d{1,4}", s):
        return None
    match = re.search(r"[-+]?\d*\.?\d+", s)
    if match:
        try:
            return float(match.group(0))
        except ValueError:
            return None
    return None


def record_to_values_dict(record: ExtractedRecord) -> Dict[str, Any]:
    """Snapshot a record's current field values for audit history."""
    return {
        "record_type": record.record_type,
        "category": record.category,
        "part_number": record.part_number,
        "cleaned_part_number": record.cleaned_part_number,
        "model": record.model,
        "serial_number": record.serial_number,
        "condition": record.condition,
        "condition_notes": record.condition_notes,
        "description": record.description,
        "quantity": record.quantity,
        "unit": record.unit,
        "vendor": record.vendor,
        "manufacturer": record.manufacturer,
        "location": record.location,
        "rack": record.rack,
        "u_position": record.u_position,
        "power_kw": record.power_kw,
        "cable_length_m": record.cable_length_m,
        "source_device": record.source_device,
        "source_port": record.source_port,
        "target_device": record.target_device,
        "target_port": record.target_port,
        "cable_type": record.cable_type,
        "review_status": record.review_status,
        "review_notes": record.review_notes,
        "review_reason": record.review_reason,
        "unmapped_columns": record.unmapped_columns or [],
        "confidence_score": record.confidence_score,
        "field_sources": record.field_sources,
    }



class ExtractionService:

    @staticmethod
    def _parse_row_dict(
        headers: List[str],
        row_values: List[Any],
        source_sheet: Optional[str],
        row_idx: int,
        col_mapping: Optional[Dict[str, str]] = None,
        col_confidence: Optional[Dict[str, float]] = None,
    ) -> Optional[Dict[str, Any]]:
        """Map a single row of tabular data into a structured record dictionary."""
        if col_mapping is None or col_confidence is None:
            col_mapping, col_confidence = map_headers(headers)

        record_data: Dict[str, Any] = {
            "source_sheet": source_sheet,
            "source_row": row_idx,
            "raw_data": {},
            "unmapped_columns": [],
            "field_sources": {},
        }

        # 1. Map columns to canonical fields or identify unmapped non-empty values
        for idx, (h, val) in enumerate(zip(headers, row_values)):
            col_label = (str(h).strip() if h is not None and str(h).strip() else f"Column {idx + 1}")
            val_str = str(val).strip() if val is not None else ""
            record_data["raw_data"][col_label] = val_str
            col_letter = col_num_to_letter(idx)

            canonical = col_mapping.get(h)
            if canonical:
                if canonical not in record_data or not record_data[canonical]:
                    record_data[canonical] = val_str
                    record_data["field_sources"][canonical] = {
                        "sheet": source_sheet or "Sheet1",
                        "row": row_idx,
                        "col": col_letter,
                        "header": col_label,
                    }
            else:
                # Unmapped column with non-empty content
                if val_str:
                    record_data["unmapped_columns"].append(col_label)

        # 2. Extract and clean part number
        part_number = record_data.get("part_number")

        if part_number:
            record_data["cleaned_part_number"] = clean_part_number(part_number)

        # 3. Model
        if not record_data.get("model") and record_data.get("cleaned_part_number"):
            record_data["model"] = record_data.get("cleaned_part_number")

        # Sync location and rack for backward compatibility
        if record_data.get("rack") and not record_data.get("location"):
            record_data["location"] = record_data.get("rack")
        elif record_data.get("location") and not record_data.get("rack"):
            record_data["rack"] = record_data.get("location")

        # 4. Strict quantity parsing (NEVER default to 1.0)
        # Issue 4: Never map date columns to quantity; serial rows have quantity null
        if record_data.get("serial_number"):
            record_data["quantity"] = None
            record_data["unit"] = None
        else:
            raw_qty = record_data.get("quantity")
            record_data["quantity"] = parse_quantity_strict(raw_qty) if raw_qty else None
            if not record_data.get("quantity"):
                record_data["unit"] = None

        # 5. Detect condition and notes from DOA/test columns
        cond, cond_notes = detect_condition(record_data["raw_data"])
        if cond:
            record_data["condition"] = cond
        if cond_notes:
            existing_notes = record_data.get("condition_notes")
            record_data["condition_notes"] = (
                f"{existing_notes}; {cond_notes}"
                if existing_notes and existing_notes != cond_notes
                else cond_notes
            )

        # 6. Description & Part number fallback
        description = record_data.get("description")
        serial_number = record_data.get("serial_number")

        # If row has no part_number, no serial_number, no description, and no quantity: skip empty row
        if not part_number and not serial_number and not description and record_data.get("quantity") is None:
            return None

        # Infer vendor / manufacturer
        if not record_data.get("vendor"):
            inferred = extract_vendor_from_text(f"{part_number or ''} {description or ''}")
            if inferred:
                record_data["vendor"] = inferred
                record_data["manufacturer"] = inferred
        elif not record_data.get("manufacturer"):
            record_data["manufacturer"] = record_data.get("vendor")

        # 7. Strict float conversions for numeric fields
        if record_data.get("power_kw"):
            record_data["power_kw"] = parse_quantity_strict(record_data.get("power_kw"))
        if record_data.get("cable_length_m"):
            record_data["cable_length_m"] = parse_quantity_strict(record_data.get("cable_length_m"))

        # 8. Compute review_reason
        reasons = []
        if record_data.get("part_number") is None:
            reasons.append("missing_value")
        elif record_data.get("quantity") is None:
            reasons.append("missing_value")

        if record_data["unmapped_columns"]:
            reasons.append("unmapped_field")

        if str(record_data.get("condition") or "").lower() in ["failed", "doa", "damaged"]:
            reasons.append("condition_failure")

        if col_confidence and any(conf < 0.9 for conf in col_confidence.values()):
            reasons.append("unclear_match")

        record_data["review_reason"] = reasons[0] if reasons else None

        # 9. Compute fine-grained confidence score (0.0 to 1.0)
        score = 1.0
        if not record_data.get("part_number"):
            score -= 0.30
        if record_data.get("quantity") is None and not record_data.get("serial_number"):
            score -= 0.20
        if record_data["unmapped_columns"]:
            score -= min(0.30, len(record_data["unmapped_columns"]) * 0.15)
        if col_confidence and any(conf < 0.9 for conf in col_confidence.values()):
            score -= 0.10
        if part_number and re.search(r"[\s]{2,}|[?<>*!]", part_number):
            score -= 0.20

        record_data["confidence_score"] = round(max(0.05, min(1.0, score)), 2)

        # 10. Copy extracted canonical fields into immutable original_values
        record_data["original_values"] = {
            "part_number": record_data.get("part_number"),
            "cleaned_part_number": record_data.get("cleaned_part_number"),
            "model": record_data.get("model"),
            "serial_number": record_data.get("serial_number"),
            "condition": record_data.get("condition"),
            "condition_notes": record_data.get("condition_notes"),
            "description": record_data.get("description"),
            "quantity": record_data.get("quantity"),
            "unit": record_data.get("unit"),
            "vendor": record_data.get("vendor"),
            "manufacturer": record_data.get("manufacturer"),
            "location": record_data.get("location"),
            "rack": record_data.get("rack"),
            "u_position": record_data.get("u_position"),
            "source_device": record_data.get("source_device"),
            "source_port": record_data.get("source_port"),
            "target_device": record_data.get("target_device"),
            "target_port": record_data.get("target_port"),
            "cable_type": record_data.get("cable_type"),
            "power_kw": record_data.get("power_kw"),
            "cable_length_m": record_data.get("cable_length_m"),
            "field_sources": record_data.get("field_sources"),
        }

        return record_data

    @staticmethod
    def _extract_from_csv(file_path: Path) -> Tuple[List[Dict[str, Any]], List[Dict[str, str]]]:
        """Parse CSV or TSV files using standard library csv module."""
        sep = "\t" if file_path.suffix.lower() == ".tsv" else ","
        records = []
        skipped_sheets = []

        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f, delimiter=sep)
            headers = None
            for row in reader:
                if any(cell.strip() for cell in row):
                    headers = [c.strip() for c in row]
                    break

            if not headers:
                return [], []

            col_mapping, col_confidence = map_headers(headers)
            has_part_num = "part_number" in col_mapping.values()
            has_desc_qty = "description" in col_mapping.values() and "quantity" in col_mapping.values()

            if not has_part_num and not has_desc_qty:
                skipped_sheets.append({
                    "name": "default",
                    "reason": "No part number or (description + quantity) column found",
                })
                return [], skipped_sheets

            for row_idx, row in enumerate(reader, start=2):
                if not any(cell.strip() for cell in row):
                    continue
                parsed = ExtractionService._parse_row_dict(
                    headers=headers,
                    row_values=row,
                    source_sheet="default",
                    row_idx=row_idx,
                    col_mapping=col_mapping,
                    col_confidence=col_confidence,
                )
                if parsed:
                    records.append(parsed)

        return records, skipped_sheets

    @staticmethod
    def _extract_serial_matrix(sheet: Any, sheet_name: str) -> List[Dict[str, Any]]:
        """
        Detect and parse multi-column serial matrix sheets (e.g. MPO Serial Numbers.xlsx).
        The part number is the column heading (row 2); creates one line per serial with
        part_number = that column's part number and serial_number = the cell.
        """
        # A serial matrix sheet must explicitly mention "serial" in header area (top 25 rows)
        max_r_scan = min(getattr(sheet, "max_row", 25) or 25, 25)
        max_c_scan = min(getattr(sheet, "max_column", 40) or 40, 40)
        has_serial_keyword = False
        for r in range(1, max_r_scan + 1):
            for c in range(1, max_c_scan + 1):
                v = sheet.cell(r, c).value
                if v and "serial" in str(v).lower():
                    has_serial_keyword = True
                    break
            if has_serial_keyword:
                break

        if not has_serial_keyword:
            return []

        invalid_pn = {"quantity", "qty", "count", "description", "item", "unit", "total", "rack", "status", "serial", "serial number", "cable part number"}
        pn_blocks = []
        max_r = min(getattr(sheet, "max_row", 20) or 20, 10)
        max_c = getattr(sheet, "max_column", 40) or 40

        for r in range(1, max_r + 1):
            for c in range(1, max_c + 1):
                v = sheet.cell(r, c).value
                if v and "part number" in str(v).lower():
                    pn_val = None
                    for c2 in range(c + 1, min(max_c + 1, c + 5)):
                        v2 = sheet.cell(r, c2).value
                        if v2 and str(v2).strip():
                            cand = str(v2).strip()
                            if cand.lower() not in invalid_pn:
                                pn_val = cand
                                break
                    if pn_val:
                        pn_blocks.append((r, c, pn_val))

        if not pn_blocks:
            return []

        matrix_records = []
        total_r = getattr(sheet, "max_row", 1200) or 1200
        for r_pn, c_pn, pn in pn_blocks:
            clean_pn = clean_part_number(pn)
            s_col = None
            s_row = None
            desc = None

            for r_scan in range(r_pn + 1, min(total_r + 1, r_pn + 12)):
                for c_scan in range(max(1, c_pn - 2), min(max_c + 1, c_pn + 5)):
                    cv = sheet.cell(r_scan, c_scan).value
                    if cv:
                        cv_str = str(cv).strip()
                        if "serial" in cv_str.lower():
                            s_col = c_scan
                            s_row = r_scan + 1
                        elif len(cv_str) > 20 and not desc:
                            desc = cv_str

            if not s_col or not s_row:
                continue

            col_letter = col_num_to_letter(s_col - 1)
            for r in range(s_row, total_r + 1):
                cell_val = sheet.cell(r, s_col).value
                if cell_val is not None:
                    val_str = str(cell_val).strip()
                    if (
                        val_str
                        and "serial" not in val_str.lower()
                        and "part" not in val_str.lower()
                        and "product" not in val_str.lower()
                    ):
                        matrix_records.append({
                            "source_sheet": sheet_name,
                            "source_row": r,
                            "category": "cable",
                            "record_type": "cable_item",
                            "part_number": pn,
                            "cleaned_part_number": clean_pn,
                            "model": clean_pn,
                            "serial_number": val_str,
                            "quantity": None,
                            "unit": None,
                            "description": desc,
                            "confidence_score": 0.95,
                            "review_reason": None,
                            "raw_data": {
                                "part_number": pn,
                                "serial_number": val_str,
                                "description": desc or "",
                            },
                            "unmapped_columns": [],
                            "field_sources": {
                                "part_number": {
                                    "sheet": sheet_name,
                                    "row": r_pn,
                                    "col": col_num_to_letter(c_pn - 1),
                                },
                                "serial_number": {
                                    "sheet": sheet_name,
                                    "row": r,
                                    "col": col_letter,
                                },
                            },
                            "original_values": {
                                "part_number": pn,
                                "cleaned_part_number": clean_pn,
                                "model": clean_pn,
                                "serial_number": val_str,
                                "quantity": None,
                                "description": desc,
                                "category": "cable",
                            },
                            "auto_clear": True,
                        })

        return matrix_records

    @staticmethod
    def _extract_from_excel(file_path: Path) -> Tuple[List[Dict[str, Any]], List[Dict[str, str]]]:
        """Parse Excel workbooks across all sheets using openpyxl with sheet skipping."""
        records = []
        skipped_sheets = []
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]

                # Problem 2 fix: Detect multi-column serial matrix sheets first
                matrix_recs = ExtractionService._extract_serial_matrix(sheet, sheet_name)
                if matrix_recs:
                    records.extend(matrix_recs)
                    continue

                rows_iter = sheet.iter_rows(values_only=True)
                headers = None
                for row in rows_iter:
                    if any(c is not None and str(c).strip() for c in row):
                        headers = [str(c).strip() if c is not None else "" for c in row]
                        break

                if not headers:
                    continue

                col_mapping, col_confidence = map_headers(headers)
                has_part_num = "part_number" in col_mapping.values()
                has_desc_qty = "description" in col_mapping.values() and "quantity" in col_mapping.values()

                # Problem 11: Require part-number column or (description + quantity), else skip sheet
                if not has_part_num and not has_desc_qty:
                    skipped_sheets.append({
                        "name": sheet_name,
                        "reason": "No part number or (description + quantity) column found",
                    })
                    continue

                for row_idx, row in enumerate(rows_iter, start=2):
                    if not any(c is not None and str(c).strip() for c in row):
                        continue
                    row_str = [str(c).strip() if c is not None else "" for c in row]
                    parsed = ExtractionService._parse_row_dict(
                        headers=headers,
                        row_values=row_str,
                        source_sheet=sheet_name,
                        row_idx=row_idx,
                        col_mapping=col_mapping,
                        col_confidence=col_confidence,
                    )
                    if parsed:
                        records.append(parsed)
            wb.close()
        except ImportError:
            # Fallback to pandas if openpyxl not directly available
            try:
                import pandas as pd
                excel_file = pd.ExcelFile(file_path)
                for sheet_name in excel_file.sheet_names:
                    df = pd.read_excel(excel_file, sheet_name=sheet_name, dtype=str)
                    headers = [str(c) for c in df.columns]
                    col_mapping, col_confidence = map_headers(headers)
                    has_part_num = "part_number" in col_mapping.values()
                    has_desc_qty = "description" in col_mapping.values() and "quantity" in col_mapping.values()

                    if not has_part_num and not has_desc_qty:
                        skipped_sheets.append({
                            "name": sheet_name,
                            "reason": "No part number or (description + quantity) column found",
                        })
                        continue

                    for row_idx, row in df.iterrows():
                        row_vals = [str(v) if pd.notna(v) else "" for v in row]
                        parsed = ExtractionService._parse_row_dict(
                            headers=headers,
                            row_values=row_vals,
                            source_sheet=sheet_name,
                            row_idx=int(row_idx) + 1,
                            col_mapping=col_mapping,
                            col_confidence=col_confidence,
                        )
                        if parsed:
                            records.append(parsed)
            except Exception as e:
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                    detail=f"Failed to parse Excel workbook: {str(e)}",
                )
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Failed to read Excel workbook: {str(e)}",
            )

        return records, skipped_sheets

    @staticmethod
    def _extract_from_pdf(file_path: Path) -> Tuple[List[Dict[str, Any]], List[Dict[str, str]]]:
        """Extract structured line items from PDF documents."""
        records = []
        skipped_sheets = []
        try:
            import importlib
            try:
                fitz = importlib.import_module("pymupdf")
            except ImportError:
                fitz = importlib.import_module("fitz")

            doc = fitz.open(file_path)
            for page_idx, page in enumerate(doc, start=1):
                tables = getattr(page, "find_tables", lambda: None)()
                if tables and tables.tables:
                    for t_idx, table in enumerate(tables.tables):
                        table_data = table.extract()
                        if table_data and len(table_data) > 1:
                            headers = [str(c).strip() if c else "" for c in table_data[0]]
                            col_mapping, col_confidence = map_headers(headers)
                            for r_idx, r in enumerate(table_data[1:], start=2):
                                row_str = [str(c).strip() if c else "" for c in r]
                                parsed = ExtractionService._parse_row_dict(
                                    headers=headers,
                                    row_values=row_str,
                                    source_sheet=f"Page_{page_idx}_Table_{t_idx + 1}",
                                    row_idx=r_idx,
                                    col_mapping=col_mapping,
                                    col_confidence=col_confidence,
                                )
                                if parsed:
                                    records.append(parsed)
                else:
                    text = page.get_text("text")
                    for line_idx, line in enumerate(text.splitlines(), start=1):
                        parts = line.strip().split()
                        if len(parts) >= 2 and any(char.isdigit() for char in line):
                            clean_pn = clean_part_number(parts[0])
                            records.append({
                                "source_sheet": f"Page_{page_idx}",
                                "source_row": line_idx,
                                "part_number": parts[0],
                                "cleaned_part_number": clean_pn,
                                "description": " ".join(parts[1:]),
                                "quantity": None,
                                "confidence_score": 0.60,
                                "review_reason": "missing_value",
                                "raw_data": {"line": line},
                                "original_values": {
                                    "part_number": parts[0],
                                    "cleaned_part_number": clean_pn,
                                    "description": " ".join(parts[1:]),
                                    "quantity": None,
                                },
                            })
            doc.close()
        except ImportError:
            clean_pn = "DOC-SUMMARY"
            records.append({
                "source_sheet": "PDF_Extraction",
                "source_row": 1,
                "part_number": clean_pn,
                "cleaned_part_number": clean_pn,
                "description": f"Extracted from {file_path.name}",
                "quantity": None,
                "confidence_score": 0.50,
                "review_reason": "missing_value",
                "raw_data": {"filename": file_path.name},
                "original_values": {
                    "part_number": clean_pn,
                    "cleaned_part_number": clean_pn,
                    "description": f"Extracted from {file_path.name}",
                    "quantity": None,
                },
            })
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Failed to parse PDF document: {str(e)}",
            )
        return records, skipped_sheets

    @staticmethod
    def extract_document_text(file_path: Path) -> str:
        """Extract plain text lines from files for LLM prompt context."""
        ext = file_path.suffix.lower()
        if ext in [".csv", ".tsv", ".txt"]:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                return f.read(50000)
        elif ext in [".xlsx", ".xls", ".xlsm"]:
            from app.parsers.excel_parser import parse_excel_file, excel_to_llm_text
            parsed = parse_excel_file(file_path)
            return excel_to_llm_text(parsed)
        elif ext == ".pdf":
            from app.parsers.pdf_parser import parse_pdf_file, pdf_to_llm_text

            parsed = parse_pdf_file(file_path)
            pages = parsed.get("pages", [])
            pages_needing_ocr = [
                page["page_number"]
                for page in pages
                if not page.get("text", "").strip()
                and not any(
                    cell.get("value", "").strip()
                    for table in page.get("tables", [])
                    for row in table.get("rows", [])
                    for cell in row.get("cells", [])
                )
            ]

            ocr_pages = []
            if pages_needing_ocr:
                try:
                    try:
                        import pymupdf as fitz
                    except ImportError:
                        import fitz
                    import pytesseract
                    from PIL import Image

                    with fitz.open(file_path) as pdf:
                        for page_number in pages_needing_ocr:
                            page = pdf[page_number - 1]
                            pixmap = page.get_pixmap(
                                matrix=fitz.Matrix(2, 2),
                                alpha=False,
                            )
                            image = Image.frombytes(
                                "RGB",
                                (pixmap.width, pixmap.height),
                                pixmap.samples,
                            )
                            page_text = pytesseract.image_to_string(image).strip()
                            if page_text:
                                ocr_pages.append(
                                    f"===== PAGE {page_number} OCR =====\n{page_text}"
                                )
                except Exception as ocr_err:
                    logger.warning("OCR processing skipped for %s: %s", file_path, ocr_err)

            embedded_text = (
                pdf_to_llm_text(parsed)
                if any(
                    page.get("text", "").strip()
                    or any(
                        cell.get("value", "").strip()
                        for table in page.get("tables", [])
                        for row in table.get("rows", [])
                        for cell in row.get("cells", [])
                    )
                    for page in pages
                )
                else ""
            )
            return "\n".join(
                part for part in (embedded_text, *ocr_pages) if part
            )[:50000]
        elif ext == ".docx":
            from app.parsers.docx_parser import parse_docx_file, docx_to_llm_text

            parsed = parse_docx_file(file_path)
            has_content = bool(parsed["paragraphs"]) or any(
                cell["text"]
                for table in parsed["tables"]
                for row in table["rows"]
                for cell in row["cells"]
            )
            return docx_to_llm_text(parsed) if has_content else ""
        elif ext in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}:
            from app.parsers.image_parser import parse_image_file, image_to_llm_text

            parsed = parse_image_file(file_path)
            return image_to_llm_text(parsed) if parsed.get("full_text", "").strip() else ""
        return ""

    @staticmethod
    def _ai_source_content(file_path: Path) -> Tuple[str, List[Dict[str, Any]]]:
        """Return model input and parser-derived text segments with source coordinates."""
        ext = file_path.suffix.lower()
        segments: List[Dict[str, Any]] = []

        if ext in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
            from app.parsers.excel_parser import (
                excel_to_llm_text,
                parse_excel_file,
            )

            parsed = parse_excel_file(file_path)
            for sheet in parsed["sheets"]:
                for row in sheet["rows"]:
                    for cell in row["cells"]:
                        if cell["value"] is None:
                            continue
                        segments.append({
                            "text": str(cell["value"]),
                            "sheet": sheet["sheet_name"],
                            "row": row["row_number"],
                            "col": cell["column_letter"],
                            "source_type": "excel_cell",
                        })
            return excel_to_llm_text(parsed), segments

        if ext == ".xls":
            import xlrd

            workbook = xlrd.open_workbook(
                file_path,
                on_demand=True,
            )
            sheets = []
            for sheet_index in range(workbook.nsheets):
                sheet = workbook.sheet_by_index(sheet_index)
                rows = []
                for row_index in range(sheet.nrows):
                    cells = []
                    row_has_values = False
                    for column_index in range(sheet.ncols):
                        value = sheet.cell_value(row_index, column_index)
                        if value == "":
                            normalized_value = None
                        elif sheet.cell_type(row_index, column_index) == xlrd.XL_CELL_DATE:
                            normalized_value = xlrd.xldate_as_datetime(
                                value,
                                workbook.datemode,
                            ).isoformat()
                        elif isinstance(value, float) and value.is_integer():
                            normalized_value = int(value)
                        else:
                            normalized_value = value
                        if normalized_value is not None:
                            row_has_values = True
                            segments.append({
                                "text": str(normalized_value),
                                "sheet": sheet.name,
                                "row": row_index + 1,
                                "col": col_num_to_letter(column_index),
                                "source_type": "excel_cell",
                            })
                        cells.append(
                            f"COLUMN {col_num_to_letter(column_index)} "
                            f"(index {column_index + 1}): {normalized_value}"
                        )
                    if row_has_values:
                        rows.append(
                            f"ROW {row_index + 1}: " + " | ".join(cells)
                        )
                sheets.append(
                    f"===== SHEET: {sheet.name} =====\n"
                    + "\n".join(rows)
                )
            workbook.release_resources()
            return "\n".join(sheets)[:50000], segments

        if ext in {".csv", ".tsv"}:
            delimiter = "\t" if ext == ".tsv" else ","
            with file_path.open("r", encoding="utf-8", errors="replace", newline="") as source_file:
                rows = list(csv.reader(source_file, delimiter=delimiter))
            for row_number, row in enumerate(rows, start=1):
                for column_number, value in enumerate(row, start=1):
                    if value:
                        segments.append({
                            "text": value,
                            "sheet": "default",
                            "row": row_number,
                            "col": col_num_to_letter(column_number - 1),
                            "source_type": "table_cell",
                        })
            return file_path.read_text(encoding="utf-8", errors="replace")[:50000], segments

        if ext == ".docx":
            from app.parsers.docx_parser import parse_docx_file, docx_to_llm_text

            parsed = parse_docx_file(file_path)
            for paragraph in parsed["paragraphs"]:
                segments.append({
                    "text": paragraph["text"],
                    "paragraph": paragraph["paragraph_number"],
                    "source_type": "paragraph",
                })
            for table in parsed["tables"]:
                for row in table["rows"]:
                    for cell in row["cells"]:
                        if cell["text"]:
                            segments.append({
                                "text": cell["text"],
                                "table": table["table_number"],
                                "row": row["row_number"],
                                "col": cell["column_position"],
                                "source_type": "table_cell",
                            })
            return docx_to_llm_text(parsed)[:50000], segments

        if ext in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}:
            from app.parsers.image_parser import parse_image_file, image_to_llm_text

            parsed = parse_image_file(file_path)
            image = parsed["image"]
            for row in parsed["rows"]:
                cells = row["cells"]
                if not cells:
                    continue
                left = min(cell["left"] for cell in cells)
                top = min(cell["top"] for cell in cells)
                right = max(cell["left"] + cell["width"] for cell in cells)
                bottom = max(cell["top"] + cell["height"] for cell in cells)
                segments.append({
                    "text": " ".join(cell["text"] for cell in cells),
                    "page": 1,
                    "bbox": [left, top, right, bottom],
                    "word_boxes": [
                        {
                            "text": cell["text"],
                            "bbox": [
                                cell["left"],
                                cell["top"],
                                cell["left"] + cell["width"],
                                cell["top"] + cell["height"],
                            ],
                        }
                        for cell in cells
                    ],
                    "page_width": image["width"],
                    "page_height": image["height"],
                    "source_type": "ocr_line",
                    "ocr_confidence": sum(
                        max(0.0, cell["confidence"]) / 100
                        for cell in cells
                    ) / len(cells),
                })
            return image_to_llm_text(parsed)[:50000], segments

        if ext == ".pdf":
            from app.parsers.pdf_parser import parse_pdf_file, pdf_to_llm_text

            parsed = parse_pdf_file(file_path)
            ocr_pages: List[str] = []
            ocr_page_sizes: Dict[int, Tuple[int, int]] = {}
            pages_needing_ocr = [
                page["page_number"]
                for page in parsed.get("pages", [])
                if not page.get("text", "").strip()
                and not any(
                    cell.get("value", "").strip()
                    for table in page.get("tables", [])
                    for row in table.get("rows", [])
                    for cell in row.get("cells", [])
                )
            ]
            ocr_words_by_page: Dict[int, List[Dict[str, Any]]] = {}
            if pages_needing_ocr:
                try:
                    try:
                        import pymupdf as fitz
                    except ImportError:
                        import fitz
                    import pytesseract
                    from PIL import Image
                    from pytesseract import Output

                    with fitz.open(file_path) as pdf:
                        for page_number in pages_needing_ocr:
                            page = pdf[page_number - 1]
                            pixmap = page.get_pixmap(
                                matrix=fitz.Matrix(2, 2),
                                alpha=False,
                            )
                            ocr_page_sizes[page_number] = (
                                pixmap.width,
                                pixmap.height,
                            )
                            image = Image.frombytes(
                                "RGB",
                                (pixmap.width, pixmap.height),
                                pixmap.samples,
                            )
                            data = pytesseract.image_to_data(
                                image,
                                output_type=Output.DICT,
                            )
                            words = []
                            for index, raw_text in enumerate(data.get("text", [])):
                                word = str(raw_text).strip()
                                if not word:
                                    continue
                                try:
                                    confidence = float(data["conf"][index])
                                except (TypeError, ValueError):
                                    confidence = 0.0
                                words.append({
                                    "text": word,
                                    "left": int(data["left"][index]),
                                    "top": int(data["top"][index]),
                                    "width": int(data["width"][index]),
                                    "height": int(data["height"][index]),
                                    "line": (
                                        data["block_num"][index],
                                        data["par_num"][index],
                                        data["line_num"][index],
                                    ),
                                    "confidence": confidence,
                                })
                            ocr_words_by_page[page_number] = words
                            ocr_text = pytesseract.image_to_string(image).strip()
                            if ocr_text:
                                ocr_pages.append(
                                    f"===== PAGE {page_number} OCR =====\n{ocr_text}"
                                )
                except Exception as ocr_err:
                    logger.warning("OCR processing skipped for %s: %s", file_path, ocr_err)

            for page in parsed["pages"]:
                page_number = page["page_number"]
                words = page.get("words", [])
                grouped_words: List[List[Dict[str, Any]]] = []
                for word in sorted(words, key=lambda item: (item["top"], item["x0"])):
                    if not grouped_words or abs(
                        word["top"] - grouped_words[-1][0]["top"]
                    ) > 3:
                        grouped_words.append([word])
                    else:
                        grouped_words[-1].append(word)
                for line_words in grouped_words:
                    segments.append({
                        "text": " ".join(word["text"] for word in line_words),
                        "page": page_number,
                        "bbox": [
                            min(word["x0"] for word in line_words),
                            min(word["top"] for word in line_words),
                            max(word["x1"] for word in line_words),
                            max(word["bottom"] for word in line_words),
                        ],
                        "word_boxes": [
                            {
                                "text": word["text"],
                                "bbox": [
                                    word["x0"],
                                    word["top"],
                                    word["x1"],
                                    word["bottom"],
                                ],
                            }
                            for word in line_words
                        ],
                        "page_width": page["width"],
                        "page_height": page["height"],
                        "source_type": "pdf_text_line",
                    })

                if page_number in ocr_words_by_page:
                    words = ocr_words_by_page[page_number]
                    grouped: Dict[Tuple[int, int, int], List[Dict[str, Any]]] = {}
                    for word in words:
                        grouped.setdefault(word["line"], []).append(word)
                    page_width = page["width"]
                    page_height = page["height"]
                    pix_width, pix_height = ocr_page_sizes[page_number]
                    for line_words in grouped.values():
                        left = min(word["left"] for word in line_words)
                        top = min(word["top"] for word in line_words)
                        right = max(word["left"] + word["width"] for word in line_words)
                        bottom = max(word["top"] + word["height"] for word in line_words)
                        segments.append({
                            "text": " ".join(word["text"] for word in line_words),
                            "page": page_number,
                            "bbox": [
                                left * page_width / pix_width,
                                top * page_height / pix_height,
                                right * page_width / pix_width,
                                bottom * page_height / pix_height,
                            ],
                            "word_boxes": [
                                {
                                    "text": word["text"],
                                    "bbox": [
                                        word["left"] * page_width / pix_width,
                                        word["top"] * page_height / pix_height,
                                        (word["left"] + word["width"])
                                        * page_width / pix_width,
                                        (word["top"] + word["height"])
                                        * page_height / pix_height,
                                    ],
                                }
                                for word in line_words
                            ],
                            "page_width": page_width,
                            "page_height": page_height,
                            "source_type": "ocr_line",
                            "ocr_confidence": sum(
                                max(0.0, word["confidence"]) / 100
                                for word in line_words
                            ) / len(line_words),
                        })

            content = pdf_to_llm_text(parsed)
            if ocr_pages:
                content = f"{content}\n\n" + "\n\n".join(ocr_pages)
            return content[:50000], segments

        if ext == ".txt":
            content = file_path.read_text(encoding="utf-8", errors="replace")[:50000]
            for line_number, line in enumerate(content.splitlines(), start=1):
                if line.strip():
                    segments.append({
                        "text": line.strip(),
                        "line": line_number,
                        "source_type": "text_line",
                    })
            return content, segments

        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"AI extraction does not support {ext or 'this file type'}.",
        )

    @staticmethod
    def _verify_ai_value(
        value: Any,
        segments: List[Dict[str, Any]],
    ) -> Optional[Dict[str, Any]]:
        """Find an exact textual or numeric value in parser-derived source segments."""
        if value is None or isinstance(value, bool):
            return None

        text_value = str(value).strip()
        if not text_value:
            return None

        numeric_value: Optional[Decimal] = None
        if isinstance(value, (int, float)):
            try:
                numeric_value = Decimal(str(value))
            except InvalidOperation:
                numeric_value = None

        for segment in segments:
            source_text = str(segment["text"])
            matches = [
                (match.start(), match.end(), match.group())
                for match in re.finditer(re.escape(text_value), source_text)
            ]
            if numeric_value is not None:
                for match in re.finditer(
                    r"(?<![\w.])[-+]?\d+(?:\.\d+)?(?![\w.])",
                    source_text,
                ):
                    try:
                        if Decimal(match.group()) == numeric_value:
                            matches.append(
                                (match.start(), match.end(), match.group())
                            )
                    except InvalidOperation:
                        continue

            for match_start, match_end, matched_text in sorted(matches):
                if match_start > 0 and re.match(
                    r"[\w./-]",
                    source_text[match_start - 1],
                ):
                    continue
                if match_end < len(source_text) and re.match(
                    r"[\w./-]",
                    source_text[match_end],
                ):
                    continue

                provenance = dict(segment)
                provenance["text"] = matched_text
                word_boxes = segment.get("word_boxes")
                if word_boxes:
                    character_offset = 0
                    selected_boxes = []
                    for word_box in word_boxes:
                        word_start = character_offset
                        word_end = word_start + len(word_box["text"])
                        if word_start < match_end and word_end > match_start:
                            selected_boxes.append(word_box["bbox"])
                        character_offset = word_end + 1
                    if selected_boxes:
                        provenance["bbox"] = [
                            min(box[0] for box in selected_boxes),
                            min(box[1] for box in selected_boxes),
                            max(box[2] for box in selected_boxes),
                            max(box[3] for box in selected_boxes),
                        ]
                provenance.pop("word_boxes", None)
                if segment.get("ocr_confidence") is not None:
                    provenance["ocr_confidence"] = round(
                        segment["ocr_confidence"],
                        3,
                    )
                return provenance
        return None

    @staticmethod
    def _ai_image_inputs(file_path: Path) -> List[Tuple[bytes, str]]:
        ext = file_path.suffix.lower()
        if ext in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}:
            if ext in {".png", ".jpg", ".jpeg", ".webp"}:
                mime_type = {
                    ".jpg": "image/jpeg",
                    ".jpeg": "image/jpeg",
                }.get(ext, f"image/{ext.lstrip('.')}")
                return [(file_path.read_bytes(), mime_type)]

            from PIL import Image

            with Image.open(file_path) as image:
                converted = image.convert("RGB")
                output = io.BytesIO()
                converted.save(output, format="PNG")
                return [(output.getvalue(), "image/png")]

        if ext == ".pdf":
            try:
                try:
                    import pymupdf as fitz
                except ImportError:
                    import fitz

                images = []
                with fitz.open(file_path) as pdf:
                    for page in pdf:
                        pixmap = page.get_pixmap(
                            matrix=fitz.Matrix(1.5, 1.5),
                            alpha=False,
                        )
                        images.append((pixmap.tobytes("png"), "image/png"))
                return images
            except Exception as img_err:
                logger.warning("Failed to render PDF images for %s: %s", file_path, img_err)
                return []

        return []

    @staticmethod
    def extract_document(
        db: Session,
        document_id: str,
        use_ai: bool = False,
        model: Optional[str] = None,
    ) -> ExtractionJobResponse:
        """Execute extraction pipeline for a specific document and store extracted records."""
        if use_ai:
            return ExtractionService.extract_document_with_ai(
                db=db,
                document_id=document_id,
                model=model,
            )

        doc = db.execute(select(Document).where(Document.id == document_id)).scalar_one_or_none()
        if not doc:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Document '{document_id}' not found.",
            )

        file_path = Path(doc.file_path)
        if not file_path.exists():
            alt_path = Path(settings.UPLOAD_DIR) / doc.project_id / doc.stored_filename
            if alt_path.exists():
                file_path = alt_path
            else:
                matches = list(Path(settings.UPLOAD_DIR).glob(f"**/{doc.stored_filename}"))
                if matches:
                    file_path = matches[0]
                else:
                    doc.status = DocumentStatus.FAILED.value
                    doc.error_message = f"File not found on storage: {doc.file_path}"
                    db.commit()
                    raise HTTPException(
                        status_code=status.HTTP_404_NOT_FOUND,
                        detail="Source file missing from disk storage.",
                    )

        doc.status = DocumentStatus.EXTRACTING.value
        db.commit()

        ext = file_path.suffix.lower()
        if ext in [".pdf", ".docx", ".doc", ".png", ".jpg", ".jpeg"]:
            from app.ai.llm import ollama_client
            if ollama_client.is_available():
                return ExtractionService.extract_document_with_ai(
                    db=db,
                    document_id=document_id,
                    model=model,
                    allow_fallback=True,
                )

        extracted_data: List[Dict[str, Any]] = []
        skipped_sheets: List[Dict[str, str]] = []

        try:
            if ext in [".csv", ".tsv"]:
                extracted_data, skipped_sheets = ExtractionService._extract_from_csv(file_path)
            elif ext in [".xlsx", ".xls", ".xlsm"]:
                extracted_data, skipped_sheets = ExtractionService._extract_from_excel(file_path)
            elif ext == ".pdf":
                extracted_data, skipped_sheets = ExtractionService._extract_from_pdf(file_path)
            else:
                extracted_data, skipped_sheets = ExtractionService._extract_from_csv(file_path)

            if len(extracted_data) == 0:
                from app.ai.llm import ollama_client
                if ollama_client.is_available():
                    return ExtractionService.extract_document_with_ai(
                        db=db,
                        document_id=document_id,
                        model=model,
                        allow_fallback=True,
                    )

            # Store skipped sheets in document metadata
            doc_meta = dict(doc.metadata_json or {})
            if skipped_sheets:
                doc_meta["skipped_sheets"] = skipped_sheets
            doc.metadata_json = doc_meta

            doc_type_mapping = {
                "bom": RecordType.BOM_ITEM.value,
                "purchase": RecordType.BOM_ITEM.value,
                "p2p": RecordType.CABLE_ITEM.value,
                "rack_layout": RecordType.RACK_ITEM.value,
                "power": RecordType.POWER_ITEM.value,
                "cable": RecordType.CABLE_ITEM.value,
            }
            default_record_type = doc_type_mapping.get(doc.document_type, RecordType.BOM_ITEM.value)

            # Problem 13: Deduplication and review status preservation on re-extract
            existing_records = db.execute(
                select(ExtractedRecord).where(
                    ExtractedRecord.document_id == document_id,
                    ExtractedRecord.is_deleted == False,
                )
            ).scalars().all()

            by_sheet_row = {
                (r.source_sheet, r.source_row): r
                for r in existing_records
                if r.source_row is not None
            }
            by_pn_ident = {
                (r.part_number, r.serial_number or r.location or r.rack): r
                for r in existing_records
                if r.part_number
            }

            matched_existing_ids = set()

            for item in extracted_data:
                # Find matching existing record
                match = by_sheet_row.get((item.get("source_sheet"), item.get("source_row")))
                if not match and item.get("part_number"):
                    ident_key = (
                        item.get("part_number"),
                        item.get("serial_number") or item.get("location") or item.get("rack"),
                    )
                    match = by_pn_ident.get(ident_key)

                if match:
                    matched_existing_ids.add(match.id)
                    # Check if core parsed values changed
                    values_changed = (
                        match.part_number != item.get("part_number")
                        or match.quantity != item.get("quantity")
                        or match.serial_number != item.get("serial_number")
                        or match.condition != item.get("condition")
                        or match.condition_notes != item.get("condition_notes")
                        or match.description != item.get("description")
                    )
                    if values_changed:
                        # Values changed: update and reset to pending_review
                        old_val = record_to_values_dict(match)
                        match.part_number = item.get("part_number")
                        match.cleaned_part_number = item.get("cleaned_part_number")
                        match.model = item.get("model")
                        match.serial_number = item.get("serial_number")
                        match.condition = item.get("condition")
                        match.condition_notes = item.get("condition_notes")
                        match.description = item.get("description")
                        match.quantity = item.get("quantity")
                        match.unit = item.get("unit")
                        match.vendor = item.get("vendor")
                        match.manufacturer = item.get("manufacturer")
                        match.location = item.get("location")
                        match.rack = item.get("rack")
                        match.u_position = item.get("u_position")
                        match.power_kw = item.get("power_kw")
                        match.cable_length_m = item.get("cable_length_m")
                        match.source_device = item.get("source_device")
                        match.target_device = item.get("target_device")
                        match.raw_data = item.get("raw_data")
                        match.unmapped_columns = item.get("unmapped_columns", [])
                        match.review_reason = item.get("review_reason")
                        match.confidence_score = item.get("confidence_score", 0.80)
                        match.review_status = ReviewStatus.PENDING_REVIEW.value
                        new_val = record_to_values_dict(match)

                        db.add(
                            RecordReviewAction(
                                id=generate_id("act_"),
                                record_id=match.id,
                                action="re_extract_changed",
                                old_values=old_val,
                                new_values=new_val,
                                reason="Values changed during re-extraction; reset to pending review",
                                reviewer_id="system",
                            )
                        )
                    else:
                        # Values unchanged: keep approvals and review state intact!
                        match.raw_data = item.get("raw_data")
                        match.unmapped_columns = item.get("unmapped_columns", [])
                else:
                    # New record
                    # Issue 1: Lines that pass every check get review_status "read_clearly" (or approved by "system")
                    is_serial = bool(item.get("serial_number"))
                    has_pn = bool(item.get("part_number") or item.get("cleaned_part_number"))
                    has_qty = is_serial or (item.get("quantity") is not None and isinstance(item.get("quantity"), (int, float)))
                    all_mapped = len(item.get("unmapped_columns") or []) == 0
                    cond = str(item.get("condition") or "").lower()
                    no_conflict = cond not in ["failed", "doa", "damaged", "fail", "broken"]
                    high_score = (item.get("confidence_score") or 0.0) >= 0.85

                    passes_all = (has_pn and has_qty and all_mapped and no_conflict and high_score) or item.get("auto_clear", False)
                    if passes_all:
                        if random.random() < 0.05:
                            status_val = ReviewStatus.PENDING_REVIEW.value
                            reason_val = "Random 5% quality spot check"
                            rev_by = None
                            rev_at = None
                        else:
                            status_val = "read_clearly"
                            reason_val = None
                            rev_by = "system"
                            rev_at = datetime.now(timezone.utc)
                    else:
                        status_val = ReviewStatus.PENDING_REVIEW.value
                        reason_val = item.get("review_reason") or "Requires manual review"
                        rev_by = None
                        rev_at = None

                    rec = ExtractedRecord(
                        id=generate_id("rec"),
                        project_id=doc.project_id,
                        document_id=doc.id,
                        record_type=item.get("record_type", default_record_type),
                        category=item.get("category"),
                        part_number=item.get("part_number"),
                        cleaned_part_number=item.get("cleaned_part_number"),
                        model=item.get("model"),
                        serial_number=item.get("serial_number"),
                        condition=item.get("condition"),
                        condition_notes=item.get("condition_notes"),
                        description=item.get("description"),
                        quantity=item.get("quantity"),
                        unit=item.get("unit"),
                        vendor=item.get("vendor"),
                        manufacturer=item.get("manufacturer"),
                        location=item.get("location"),
                        rack=item.get("rack"),
                        u_position=item.get("u_position"),
                        power_kw=item.get("power_kw"),
                        cable_length_m=item.get("cable_length_m"),
                        source_device=item.get("source_device"),
                        source_port=item.get("source_port"),
                        target_device=item.get("target_device"),
                        target_port=item.get("target_port"),
                        cable_type=item.get("cable_type"),
                        confidence_score=item.get("confidence_score", 0.80),
                        raw_data=item.get("raw_data"),
                        unmapped_columns=item.get("unmapped_columns", []),
                        review_reason=reason_val,
                        original_values=item.get("original_values"),
                        extraction_method=item.get("extraction_method", "table"),
                        model_used=item.get("model_used"),
                        source_sheet=item.get("source_sheet"),
                        source_row=item.get("source_row"),
                        field_sources=item.get("field_sources"),
                        review_status=status_val,
                        reviewed_by=rev_by,
                        reviewed_at=rev_at,
                        is_manually_added=False,
                        is_deleted=False,
                    )
                    db.add(rec)

            # Mark missing records as removed
            for r in existing_records:
                if r.id not in matched_existing_ids:
                    r.is_deleted = True

            doc.status = DocumentStatus.EXTRACTED.value
            doc.error_message = None
            db.commit()

            # Record audit log for extraction
            try:
                from app.api.audit import record_audit_log
                record_audit_log(
                    db=db,
                    action="extract",
                    resource_type="document",
                    resource_id=doc.id,
                    project_id=doc.project_id,
                    details={
                        "filename": doc.filename,
                        "records_extracted": len(extracted_data),
                    },
                )
            except Exception:
                pass

            return ExtractionJobResponse(
                document_id=doc.id,
                document_name=doc.filename,
                status="extracted",
                records_extracted=len(extracted_data),
                message=f"Successfully extracted {len(extracted_data)} records.",
                model_used=None,
                fallback=False,
            )

        except HTTPException:
            raise
        except Exception as e:
            doc.status = DocumentStatus.FAILED.value
            doc.error_message = f"Extraction failure: {str(e)}"
            db.commit()
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Extraction failed: {str(e)}",
            )


    @staticmethod
    def extract_document_with_ai(
        db: Session,
        document_id: str,
        model: Optional[str] = None,
        allow_fallback: bool = False,
    ) -> ExtractionJobResponse:
        """Run AI LLM extraction on a document."""
        from app.ai.llm import llm_client

        target_model = model or settings.LLM_MODEL
        doc = db.execute(select(Document).where(Document.id == document_id)).scalar_one_or_none()
        if not doc:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Document '{document_id}' not found.",
            )

        file_path = Path(doc.file_path)
        if not file_path.exists():
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Source file missing from disk storage.",
            )
        if not llm_client.is_available(target_model):
            if allow_fallback:
                return ExtractionService.extract_document(
                    db=db,
                    document_id=document_id,
                    use_ai=False,
                )
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="AI engine is not available",
            )

        from app.ai.chains import extract_records_with_llm
        doc.status = DocumentStatus.EXTRACTING.value
        db.commit()

        try:
            text_content, source_segments = ExtractionService._ai_source_content(
                file_path
            )
            if not text_content.strip():
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                    detail="No readable text was found in the document for AI extraction.",
                )

            ai_result = extract_records_with_llm(
                text_content=text_content,
                document_type=doc.document_type,
                model=target_model,
                source_file_name=doc.filename,
                source_file_type=file_path.suffix.lower(),
                image_inputs=ExtractionService._ai_image_inputs(file_path),
            )
            ai_records = [r.model_dump() for r in ai_result.records] if hasattr(ai_result, "records") else ai_result
            if not ai_records:
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                    detail="AI extraction returned no records.",
                )

            for item in ai_records:
                source_fields = {
                    "equipment_name": "equipment_name",
                    "part_number": "part_number",
                    "serial_number": "serial_number",
                    "description": "description",
                    "model": "model",
                    "manufacturer": "manufacturer",
                    "vendor": "vendor",
                    "condition": "condition",
                    "condition_notes": "condition_notes",
                    "quantity": "quantity",
                    "unit": "unit",
                    "location": "location",
                    "rack": "rack",
                    "rack_id": "rack_id",
                    "u_position": "u_position",
                    "source_device": "source_device",
                    "source_port": "source_port",
                    "target_device": "target_device",
                    "target_port": "target_port",
                    "cable_type": "cable_type",
                    "length": "cable_length",
                    "power_watts": "power_watts",
                    "dimensions": "dimensions",
                }
                verified_values: Dict[str, Any] = {}
                field_sources: Dict[str, Dict[str, Any]] = {}
                unverified_fields: List[str] = []
                candidate_count = 0
                for ai_field, record_field in source_fields.items():
                    value = item.get(ai_field)
                    if value is None or value == "":
                        continue
                    candidate_count += 1
                    provenance = ExtractionService._verify_ai_value(
                        value,
                        source_segments,
                    )
                    if provenance is None:
                        unverified_fields.append(ai_field)
                        continue
                    verified_values[record_field] = value
                    field_sources[record_field] = provenance

                pn = verified_values.get("part_number")
                clean_pn = clean_part_number(pn) if pn else None
                verified_count = len(verified_values)
                if clean_pn and "part_number" in field_sources:
                    field_sources["cleaned_part_number"] = dict(
                        field_sources["part_number"]
                    )
                qty = parse_quantity_strict(
                    verified_values.get("quantity")
                )
                confidence = (
                    verified_count / candidate_count
                    if candidate_count
                    else 0.0
                )
                review_reason = (
                    "not_verbatim" if unverified_fields else None
                )
                orig_dict = {
                    key: verified_values.get(key)
                    for key in (
                        "part_number",
                        "equipment_name",
                        "serial_number",
                        "description",
                        "model",
                        "quantity",
                        "unit",
                        "vendor",
                        "manufacturer",
                        "condition",
                        "condition_notes",
                        "location",
                        "rack",
                        "u_position",
                        "source_device",
                        "source_port",
                        "target_device",
                        "target_port",
                        "cable_type",
                        "cable_length",
                        "power_watts",
                        "dimensions",
                    )
                }
                orig_dict["cleaned_part_number"] = clean_pn
                primary_source = (
                    field_sources.get("part_number")
                    or field_sources.get("description")
                    or next(iter(field_sources.values()), {})
                )

                rec = ExtractedRecord(
                    id=generate_id("rec"),
                    project_id=doc.project_id,
                    document_id=doc.id,
                    record_type=(
                        "cable_item"
                        if doc.document_type in {"p2p", "cable"}
                        else "rack_item"
                        if doc.document_type == "rack_layout"
                        else "power_item"
                        if doc.document_type == "power"
                        else "bom_item"
                    ),
                    category=doc.document_type,
                    part_number=pn,
                    serial_number=verified_values.get("serial_number"),
                    cleaned_part_number=clean_pn,
                    model=verified_values.get("model"),
                    description=verified_values.get("description") or verified_values.get("equipment_name"),
                    quantity=qty,
                    unit=verified_values.get("unit"),
                    vendor=verified_values.get("vendor"),
                    manufacturer=verified_values.get("manufacturer"),
                    condition=verified_values.get("condition"),
                    condition_notes=verified_values.get("condition_notes"),
                    location=verified_values.get("location"),
                    rack=verified_values.get("rack") or verified_values.get("rack_id"),
                    u_position=verified_values.get("u_position"),
                    source_device=verified_values.get("source_device"),
                    source_port=verified_values.get("source_port"),
                    target_device=verified_values.get("target_device"),
                    target_port=verified_values.get("target_port"),
                    cable_type=verified_values.get("cable_type"),
                    confidence_score=confidence,
                    raw_data={
                        "unverified_fields": unverified_fields,
                        "equipment_name": verified_values.get("equipment_name"),
                        "rack_id": verified_values.get("rack_id"),
                        "dimensions": verified_values.get("dimensions"),
                        "power_watts": verified_values.get("power_watts"),
                        "cable_length": verified_values.get("cable_length"),
                    },
                    original_values=orig_dict,
                    extraction_method="ai",
                    model_used=target_model,
                    review_reason=review_reason,
                    source_sheet=primary_source.get("sheet") or (
                        f"Page_{primary_source['page']}"
                        if primary_source.get("page") is not None
                        else None
                    ),
                    source_row=primary_source.get("row"),
                    field_sources=field_sources,
                    review_status=ReviewStatus.PENDING_REVIEW.value,
                    is_manually_added=False,
                    is_deleted=False,
                )
                db.add(rec)

            doc.status = DocumentStatus.EXTRACTED.value
            doc.error_message = None
            db.commit()

            # Record audit log for AI extraction
            try:
                from app.api.audit import record_audit_log
                record_audit_log(
                    db=db,
                    action="extract",
                    resource_type="document",
                    resource_id=doc.id,
                    project_id=doc.project_id,
                    details={
                        "filename": doc.filename,
                        "records_extracted": len(ai_records),
                        "extraction_method": "ai",
                        "model_used": target_model,
                    },
                )
            except Exception:
                pass

            return ExtractionJobResponse(
                document_id=doc.id,
                document_name=doc.filename,
                status="extracted",
                records_extracted=len(ai_records),
                model_used=target_model,
                message=f"Successfully extracted {len(ai_records)} records using {target_model}.",
                fallback=False,
            )
        except HTTPException as exc:
            db.rollback()
            doc = db.get(Document, document_id)
            if doc:
                doc.status = DocumentStatus.FAILED.value
                doc.error_message = str(exc.detail)
            db.commit()
            raise
        except Exception as e:
            db.rollback()
            doc = db.get(Document, document_id)
            if doc:
                doc.status = DocumentStatus.FAILED.value
                doc.error_message = f"AI extraction error: {str(e)}"
            db.commit()
            if allow_fallback:
                try:
                    return ExtractionService.extract_document(
                        db=db,
                        document_id=document_id,
                        use_ai=False,
                    )
                except Exception:
                    pass
            if not llm_client.is_available(target_model):
                raise HTTPException(
                    status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                    detail="AI engine is not available",
                ) from e
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"AI extraction failed: {str(e)}",
            )

    @staticmethod
    def queue_document_ai_extraction(
        db: Session,
        document_id: str,
        background_tasks: BackgroundTasks,
        model: Optional[str] = None,
    ) -> ExtractionJobResponse:
        from app.ai.llm import llm_client

        target_model = model or settings.LLM_MODEL
        doc = db.execute(
            select(Document).where(Document.id == document_id)
        ).scalar_one_or_none()
        if not doc:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Document '{document_id}' not found.",
            )
        if not Path(doc.file_path).exists():
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Source file missing from disk storage.",
            )
        if not llm_client.is_available(target_model):
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="AI engine is not available",
            )

        doc.status = DocumentStatus.EXTRACTING.value
        doc.error_message = None
        db.commit()
        background_tasks.add_task(
            ExtractionService.run_background_ai_extraction,
            document_id,
            target_model,
        )
        return ExtractionJobResponse(
            document_id=doc.id,
            document_name=doc.filename,
            status="queued",
            records_extracted=0,
            message="AI extraction queued.",
            model_used=target_model,
            fallback=False,
        )

    @staticmethod
    def run_background_ai_extraction(
        document_id: str,
        model: Optional[str] = None,
    ) -> None:
        from app.core.database import SessionLocal

        with SessionLocal() as db:
            try:
                ExtractionService.extract_document_with_ai(
                    db,
                    document_id,
                    model=model,
                )
            except HTTPException as exc:
                db.rollback()
                doc = db.get(Document, document_id)
                if doc:
                    doc.status = DocumentStatus.FAILED.value
                    doc.error_message = str(exc.detail)
                    db.commit()
            except Exception as exc:
                db.rollback()
                doc = db.get(Document, document_id)
                if doc:
                    doc.status = DocumentStatus.FAILED.value
                    doc.error_message = f"AI extraction failed: {exc}"
                    db.commit()

    @staticmethod
    def extract_raw_snippet_with_ai(
        text: str,
        document_type: str = "bom",
        model: Optional[str] = None,
    ) -> List[Dict[str, Any]]:
        """Run AI extraction directly on an arbitrary text or table snippet."""
        from app.ai.chains import extract_records_with_llm
        from app.core.config import settings

        target_model = model or settings.LLM_MODEL
        try:
            return extract_records_with_llm(
                text_content=text,
                document_type=document_type,
                model=target_model,
            )
        except ConnectionError as ce:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail=str(ce),
            )
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"AI extraction error: {str(e)}",
            )

    @staticmethod
    def run_background_extraction(document_id: str) -> None:
        """Background task worker for single document extraction."""
        from app.core.database import SessionLocal
        with SessionLocal() as db:
            try:
                ExtractionService.extract_document(db, document_id)
            except Exception as e:
                db.rollback()
                doc = db.get(Document, document_id)
                if doc:
                    doc.status = DocumentStatus.FAILED.value
                    doc.error_message = f"Extraction failed: {str(e)}"
                    db.commit()

    @staticmethod
    def queue_extract_all_project_documents(
        db: Session,
        project_id: str,
        background_tasks: BackgroundTasks,
    ) -> List[ExtractionJobResponse]:
        """Problem 12: Queue background extraction and return 202 immediately."""
        project = db.execute(select(Project).where(Project.id == project_id)).scalar_one_or_none()
        if not project:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Project '{project_id}' not found.",
            )

        docs = db.execute(
            select(Document).where(Document.project_id == project_id)
        ).scalars().all()

        queued_list = []
        for d in docs:
            d.status = DocumentStatus.EXTRACTING.value
            background_tasks.add_task(ExtractionService.run_background_extraction, d.id)
            queued_list.append(
                ExtractionJobResponse(
                    document_id=d.id,
                    document_name=d.filename,
                    status="queued",
                    records_extracted=0,
                    message=f"Document '{d.filename}' queued for extraction.",
                )
            )

        db.commit()
        return queued_list

    @staticmethod
    def extract_all_project_documents(db: Session, project_id: str) -> List[ExtractionJobResponse]:
        """Synchronous version for test backward compatibility."""
        project = db.execute(select(Project).where(Project.id == project_id)).scalar_one_or_none()
        if not project:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Project '{project_id}' not found.",
            )

        docs = db.execute(
            select(Document).where(Document.project_id == project_id)
        ).scalars().all()

        results = []
        for d in docs:
            try:
                res = ExtractionService.extract_document(db, d.id)
                results.append(res)
            except HTTPException as he:
                results.append(
                    ExtractionJobResponse(
                        document_id=d.id,
                        document_name=d.filename,
                        status="failed",
                        records_extracted=0,
                        message=str(he.detail),
                    )
                )
        return results

    @staticmethod
    def list_records(
        db: Session,
        project_id: str,
        document_id: Optional[str] = None,
        review_status: Optional[str] = None,
        record_type: Optional[str] = None,
        vendor: Optional[str] = None,
        search: Optional[str] = None,
        skip: int = 0,
        limit: int = 100,
    ) -> Tuple[int, List[ExtractedRecord]]:
        """Query active extracted records with full filtering, search, and pagination."""
        query = select(ExtractedRecord).where(
            ExtractedRecord.project_id == project_id,
            ExtractedRecord.is_deleted == False,
        )

        if document_id:
            query = query.where(ExtractedRecord.document_id == document_id)
        if review_status:
            query = query.where(ExtractedRecord.review_status == review_status)
        if record_type:
            query = query.where(ExtractedRecord.record_type == record_type)
        if vendor:
            query = query.where(func.lower(ExtractedRecord.vendor) == vendor.lower())
        if search:
            search_filter = f"%{search.lower()}%"
            query = query.where(
                or_(
                    func.lower(ExtractedRecord.part_number).like(search_filter),
                    func.lower(ExtractedRecord.cleaned_part_number).like(search_filter),
                    func.lower(ExtractedRecord.description).like(search_filter),
                    func.lower(ExtractedRecord.vendor).like(search_filter),
                    func.lower(ExtractedRecord.location).like(search_filter),
                    func.lower(ExtractedRecord.rack).like(search_filter),
                )
            )

        total = db.execute(select(func.count()).select_from(query.subquery())).scalar() or 0
        records = db.execute(
            query.order_by(ExtractedRecord.created_at.asc()).offset(skip).limit(limit)
        ).scalars().all()

        return total, records

    @staticmethod
    def get_record(db: Session, record_id: str) -> ExtractedRecord:
        """Fetch a single extracted record by ID."""
        record = db.execute(
            select(ExtractedRecord).where(
                ExtractedRecord.id == record_id,
                ExtractedRecord.is_deleted == False,
            )
        ).scalar_one_or_none()
        if not record:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Extracted record '{record_id}' not found.",
            )
        return record

    @staticmethod
    def update_record(
        db: Session,
        record_id: str,
        update_data: ExtractedRecordUpdate,
        reviewer_email: str,
    ) -> ExtractedRecord:
        """Problem 9: Update fields, log to audit trail, reset approved record to pending_review."""
        record = ExtractionService.get_record(db, record_id)
        old_state = record_to_values_dict(record)

        data = update_data.model_dump(exclude_unset=True)
        review_notes = data.pop("review_notes", None)

        for key, val in data.items():
            if hasattr(record, key):
                setattr(record, key, val)

        # Re-compute cleaned part number if part number was updated
        if "part_number" in data and record.part_number:
            record.cleaned_part_number = clean_part_number(record.part_number)

        # Problem 9: PATCH on approved/modified record sets status back to pending_review
        # (or requires review_notes and sets modified)
        if review_notes and len(review_notes.strip()) >= 3:
            record.review_status = ReviewStatus.MODIFIED.value
        elif record.review_status in [ReviewStatus.APPROVED.value, ReviewStatus.MODIFIED.value]:
            record.review_status = ReviewStatus.PENDING_REVIEW.value

        record.reviewed_by = reviewer_email
        record.reviewed_at = datetime.now(timezone.utc)
        if review_notes:
            record.review_notes = review_notes

        new_state = record_to_values_dict(record)

        audit = RecordReviewAction(
            id=generate_id("act_"),
            record_id=record.id,
            action="patch",
            old_values=old_state,
            new_values=new_state,
            reason=review_notes or "Record patched via API",
            reviewer_id=reviewer_email,
        )
        db.add(audit)
        db.commit()
        db.refresh(record)
        return record

    @staticmethod
    def review_record(
        db: Session,
        record_id: str,
        review_data: ExtractedRecordReview,
        reviewer_email: str,
    ) -> ExtractedRecord:
        """Problem 8: Human review decision, audit trail, and original values restoration on undo."""
        record = ExtractionService.get_record(db, record_id)
        old_state = record_to_values_dict(record)

        action_type = "review"
        if review_data.status == ReviewStatus.PENDING_REVIEW:
            # Undo / reset action: restore original values
            action_type = "undo"
            if record.original_values:
                for k, v in record.original_values.items():
                    if hasattr(record, k):
                        setattr(record, k, v)
            record.review_status = ReviewStatus.PENDING_REVIEW.value
            record.reviewed_by = None
            record.reviewed_at = None
            record.review_notes = review_data.review_notes
        else:
            if review_data.corrections:
                corrections = review_data.corrections.model_dump(exclude_unset=True)
                for k, v in corrections.items():
                    if hasattr(record, k):
                        setattr(record, k, v)
                if "part_number" in corrections and record.part_number:
                    record.cleaned_part_number = clean_part_number(record.part_number)

            record.review_status = review_data.status.value
            record.review_notes = review_data.review_notes
            record.reviewed_by = reviewer_email
            record.reviewed_at = datetime.now(timezone.utc)

        new_state = record_to_values_dict(record)

        audit = RecordReviewAction(
            id=generate_id("act_"),
            record_id=record.id,
            action=action_type,
            old_values=old_state,
            new_values=new_state,
            reason=review_data.review_notes,
            reviewer_id=reviewer_email,
        )
        db.add(audit)
        db.commit()
        db.refresh(record)
        return record

    @staticmethod
    def batch_review(
        db: Session,
        project_id: str,
        request: BatchReviewRequest,
        reviewer_email: str,
    ) -> Tuple[int, List[Dict[str, str]]]:
        """Problem 20: Bulk review with min_confidence validation, refusal of unvetted records, and audit."""
        status_value = (
            ReviewStatus.APPROVED.value
            if request.action == BatchReviewAction.APPROVE
            else ReviewStatus.REJECTED.value
        )
        now = datetime.now(timezone.utc)

        records = db.execute(
            select(ExtractedRecord).where(
                ExtractedRecord.project_id == project_id,
                ExtractedRecord.id.in_(request.record_ids),
                ExtractedRecord.is_deleted == False,
            )
        ).scalars().all()

        updated_count = 0
        skipped: List[Dict[str, str]] = []

        found_map = {r.id: r for r in records}
        # Issue 9: List every skipped ID with its reason if not in found records
        for rec_id in request.record_ids:
            if rec_id not in found_map:
                skipped.append({
                    "id": rec_id,
                    "reason": "Record ID not found in project or is already deleted",
                })

        for r in records:
            if request.action == BatchReviewAction.APPROVE and request.min_confidence is not None:
                if r.confidence_score < request.min_confidence:
                    skipped.append({
                        "id": r.id,
                        "reason": f"Confidence score ({r.confidence_score}) is below minimum required {request.min_confidence}",
                    })
                    continue
                if r.review_reason:
                    skipped.append({
                        "id": r.id,
                        "reason": f"Record requires individual human review: flagged with '{r.review_reason}'",
                    })
                    continue
                if not r.part_number:
                    skipped.append({
                        "id": r.id,
                        "reason": "Missing part number; cannot be batch-approved",
                    })
                    continue

            old_state = record_to_values_dict(r)
            r.review_status = status_value
            r.reviewed_by = reviewer_email
            r.reviewed_at = now
            if request.review_notes:
                r.review_notes = request.review_notes
            new_state = record_to_values_dict(r)

            db.add(
                RecordReviewAction(
                    id=generate_id("act_"),
                    record_id=r.id,
                    action="batch_review",
                    old_values=old_state,
                    new_values=new_state,
                    reason=request.review_notes or f"Batch {request.action.value}",
                    reviewer_id=reviewer_email,
                )
            )
            updated_count += 1

        db.commit()

        # Audit log for batch review
        try:
            from app.api.audit import record_audit_log
            record_audit_log(
                db=db,
                action="review",
                resource_type="record",
                project_id=project_id,
                details={
                    "batch": True,
                    "action": request.action.value,
                    "updated_count": updated_count,
                    "skipped_count": len(skipped),
                    "skipped": skipped,
                },
            )
        except Exception:
            pass

        return updated_count, skipped

    @staticmethod
    def get_record_evidence(db: Session, record_id: str) -> Dict[str, Any]:
        """Problem 18: Retrieve highlighted cell or image coordinate evidence for a record."""
        from app.core.security import create_evidence_token
        record = ExtractionService.get_record(db, record_id)
        doc = db.get(Document, record.document_id)
        if not doc:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Associated document not found")

        file_path = Path(doc.file_path)
        ext = file_path.suffix.lower()
        field_sources = record.field_sources or {}
        primary_source = (
            field_sources.get("part_number")
            or field_sources.get("description")
            or next(iter(field_sources.values()), {})
        )

        # Excel evidence
        if ext in [".xlsx", ".xls", ".xlsm"]:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
                source_sheet = primary_source.get("sheet") or record.source_sheet
                sheet_name = source_sheet if source_sheet in wb.sheetnames else wb.sheetnames[0]
                sheet = wb[sheet_name]
                all_rows = list(sheet.iter_rows(values_only=True))
                wb.close()

                target_row = primary_source.get("row") or record.source_row
                if target_row is None and record.extraction_method != "ai":
                    target_row = 2
                context_row = target_row or 1
                start_row = max(1, context_row - 5)
                end_row = min(len(all_rows), context_row + 5)

                headers = [str(c) if c is not None else "" for c in all_rows[0]] if all_rows else []
                sliced_rows = [
                    [str(c) if c is not None else "" for c in row]
                    for row in all_rows[start_row - 1 : end_row]
                ]

                highlight_col = primary_source.get("col")
                if highlight_col is None and record.extraction_method != "ai":
                    highlight_col = "A"

                return {
                    "record_id": record.id,
                    "document_id": doc.id,
                    "document_name": doc.filename,
                    "evidence_type": "excel",
                    "evidence": {
                        "type": "excel",
                        "sheet": sheet_name,
                        "headers": headers,
                        "startRow": start_row,
                        "rows": sliced_rows,
                        "highlight": {
                            "row": target_row,
                            "col": highlight_col,
                            "text": primary_source.get("text"),
                        },
                        "fieldSources": field_sources,
                    },
                }
            except Exception as exc:
                if record.extraction_method == "ai":
                    raise HTTPException(
                        status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                        detail="Unable to load source evidence for the AI-extracted record.",
                    ) from exc

        # CSV / TSV evidence
        if ext in [".csv", ".tsv"]:
            try:
                sep = "\t" if ext == ".tsv" else ","
                with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                    reader = list(csv.reader(f, delimiter=sep))
                headers = reader[0] if reader else []
                target_row = primary_source.get("row") or record.source_row
                if target_row is None and record.extraction_method != "ai":
                    target_row = 2
                start_row = max(1, (target_row or 1) - 5)
                end_row = min(len(reader), (target_row or 1) + 5)
                sliced_rows = reader[start_row - 1 : end_row]

                highlight_col = primary_source.get("col")
                if highlight_col is None and record.extraction_method != "ai":
                    highlight_col = "A"

                return {
                    "record_id": record.id,
                    "document_id": doc.id,
                    "document_name": doc.filename,
                    "evidence_type": "excel",
                    "evidence": {
                        "type": "excel",
                        "sheet": record.source_sheet or "default",
                        "headers": headers,
                        "startRow": start_row,
                        "rows": sliced_rows,
                        "highlight": {
                            "row": target_row,
                            "col": highlight_col,
                            "text": primary_source.get("text"),
                        },
                        "fieldSources": field_sources,
                    },
                }
            except Exception as exc:
                if record.extraction_method == "ai":
                    raise HTTPException(
                        status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                        detail="Unable to load source evidence for the AI-extracted record.",
                    ) from exc

        if ext == ".docx":
            return {
                "record_id": record.id,
                "document_id": doc.id,
                "document_name": doc.filename,
                "evidence_type": "text",
                "evidence": {
                    "type": "text",
                    "source": primary_source,
                    "fieldSources": field_sources,
                },
            }

        token = create_evidence_token(doc.id)
        bbox = primary_source.get("bbox")
        page_width = primary_source.get("page_width")
        page_height = primary_source.get("page_height")
        normalized_bbox = None
        if bbox and page_width and page_height:
            normalized_bbox = [
                bbox[0] / page_width,
                bbox[1] / page_height,
                bbox[2] / page_width,
                bbox[3] / page_height,
            ]
        return {
            "record_id": record.id,
            "document_id": doc.id,
            "document_name": doc.filename,
            "evidence_type": "image",
            "evidence": {
                "type": "image",
                "imageUrl": f"/api/v1/documents/{doc.id}/evidence-image?token={token}",
                "width": page_width,
                "height": page_height,
                "page": primary_source.get("page"),
                "boundingBox": normalized_bbox,
                "ocrText": primary_source.get("text"),
                "confidence": record.confidence_score,
                "fieldSources": field_sources,
            },
        }

    @staticmethod
    def create_manual_record(
        db: Session,
        document_id: str,
        create_data: ExtractedRecordCreate,
        reviewer_email: str,
    ) -> ExtractedRecord:
        """Manually insert an extracted record and log creation audit."""
        doc = db.execute(select(Document).where(Document.id == document_id)).scalar_one_or_none()
        if not doc:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Document '{document_id}' not found.",
            )

        data = create_data.model_dump()
        review_notes = data.pop("review_notes", "Manually created by human reviewer")

        pn = data.get("part_number")
        clean_pn = clean_part_number(pn) if pn else None
        if clean_pn and not data.get("cleaned_part_number"):
            data["cleaned_part_number"] = clean_pn
        if clean_pn and not data.get("model"):
            data["model"] = clean_pn

        record = ExtractedRecord(
            id=generate_id("rec"),
            project_id=doc.project_id,
            document_id=doc.id,
            confidence_score=1.0,
            review_status=ReviewStatus.APPROVED.value,
            reviewed_by=reviewer_email,
            reviewed_at=datetime.now(timezone.utc),
            review_notes=review_notes,
            is_manually_added=True,
            is_deleted=False,
            original_values=dict(data),
            **data,
        )
        db.add(record)
        db.commit()
        db.refresh(record)

        db.add(
            RecordReviewAction(
                id=generate_id("act_"),
                record_id=record.id,
                action="create_manual",
                old_values=None,
                new_values=record_to_values_dict(record),
                reason=review_notes,
                reviewer_id=reviewer_email,
            )
        )
        db.commit()
        return record

    @staticmethod
    def delete_record(db: Session, record_id: str, reviewer_email: Optional[str] = None) -> bool:
        """Soft delete an extracted record and record audit action."""
        record = ExtractionService.get_record(db, record_id)
        old_state = record_to_values_dict(record)

        record.is_deleted = True
        audit = RecordReviewAction(
            id=generate_id("act_"),
            record_id=record.id,
            action="delete",
            old_values=old_state,
            new_values=None,
            reason="Soft deleted via API",
            reviewer_id=reviewer_email,
        )
        db.add(audit)
        db.commit()
        return True

    @staticmethod
    def get_record_history(db: Session, record_id: str) -> List[RecordReviewAction]:
        """Problem 8: Retrieve audit history actions for a record."""
        # Check record exists
        ExtractionService.get_record(db, record_id)
        stmt = (
            select(RecordReviewAction)
            .where(RecordReviewAction.record_id == record_id)
            .order_by(RecordReviewAction.created_at.desc())
        )
        return db.execute(stmt).scalars().all()

    @staticmethod
    def get_document_extraction_status(db: Session, document_id: str) -> DocumentExtractionStatusResponse:
        """Get extraction progress and review breakdown for a document."""
        doc = db.execute(select(Document).where(Document.id == document_id)).scalar_one_or_none()
        if not doc:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Document '{document_id}' not found.",
            )
        records = db.execute(
            select(ExtractedRecord).where(
                ExtractedRecord.document_id == document_id,
                ExtractedRecord.is_deleted == False,
            )
        ).scalars().all()

        counts = {
            "pending_review": sum(1 for r in records if r.review_status == ReviewStatus.PENDING_REVIEW.value),
            "approved": sum(1 for r in records if r.review_status == ReviewStatus.APPROVED.value),
            "modified": sum(1 for r in records if r.review_status == ReviewStatus.MODIFIED.value),
            "rejected": sum(1 for r in records if r.review_status == ReviewStatus.REJECTED.value),
        }

        return DocumentExtractionStatusResponse(
            document_id=doc.id,
            filename=doc.filename,
            document_type=doc.document_type,
            status=doc.status,
            total_records_extracted=len(records),
            review_status_counts=counts,
            error_message=doc.error_message,
        )

    @staticmethod
    def get_review_summary(db: Session, project_id: str) -> ReviewSummaryResponse:
        """Problem 21 & Issue 12: Compute review summary; return 404 for deleted or non-existent projects."""
        project = db.get(Project, project_id)
        if not project or getattr(project, "is_deleted", False):
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Project '{project_id}' not found.",
            )

        from app.models.vendor import Vendor
        records = db.execute(
            select(ExtractedRecord).where(
                ExtractedRecord.project_id == project_id,
                ExtractedRecord.is_deleted == False,
            )
        ).scalars().all()

        total = len(records)
        pending = sum(1 for r in records if r.review_status == ReviewStatus.PENDING_REVIEW.value)
        approved = sum(1 for r in records if r.review_status == ReviewStatus.APPROVED.value)
        modified = sum(1 for r in records if r.review_status == ReviewStatus.MODIFIED.value)
        rejected = sum(1 for r in records if r.review_status == ReviewStatus.REJECTED.value)

        reviewed_count = approved + modified + rejected
        completion_pct = round((reviewed_count / total * 100), 1) if total > 0 else 0.0
        avg_conf = (
            round(sum(r.confidence_score for r in records) / total, 2)
            if total > 0
            else 0.0
        )

        docs = db.execute(select(Document).where(Document.project_id == project_id)).scalars().all()
        doc_vendor_map = {}
        vendors = {v.id: v.name for v in db.execute(select(Vendor).where(Vendor.project_id == project_id)).scalars().all()}
        for d in docs:
            if d.vendor_id and d.vendor_id in vendors:
                doc_vendor_map[d.id] = vendors[d.vendor_id]
            elif d.source == "plan":
                doc_vendor_map[d.id] = "Plan (Purview)"
            else:
                doc_vendor_map[d.id] = "Unspecified Supplier"

        by_type: Dict[str, int] = {}
        by_vendor: Dict[str, int] = {}
        by_manufacturer: Dict[str, int] = {}
        for r in records:
            by_type[r.record_type] = by_type.get(r.record_type, 0) + 1
            # Supplier name from document / vendor
            supplier_name = doc_vendor_map.get(r.document_id, r.vendor or "Unspecified Supplier")
            by_vendor[supplier_name] = by_vendor.get(supplier_name, 0) + 1
            # Brand / maker
            mfg_name = r.manufacturer or r.vendor or "Unspecified Manufacturer"
            by_manufacturer[mfg_name] = by_manufacturer.get(mfg_name, 0) + 1

        return ReviewSummaryResponse(
            project_id=project_id,
            total_records=total,
            pending_review_count=pending,
            approved_count=approved,
            modified_count=modified,
            rejected_count=rejected,
            completion_percentage=completion_pct,
            average_confidence=avg_conf,
            by_record_type=by_type,
            by_vendor=by_vendor,
            by_manufacturer=by_manufacturer,
        )
