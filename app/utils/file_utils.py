import hashlib
import os
import re
from pathlib import Path
from typing import Optional, Tuple, Union

from fastapi import HTTPException, UploadFile, status

from app.schemas.document import DocumentType

ALLOWED_EXTENSIONS = {
    ".xlsx",
    ".xls",
    ".xlsm",
    ".csv",
    ".tsv",
    ".pdf",
    ".docx",
    ".doc",
    ".png",
    ".jpg",
    ".jpeg",
    ".tiff",
    ".zip",
}

EXTENSION_MIME_MAP = {
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xls": "application/vnd.ms-excel",
    ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
    ".csv": "text/csv",
    ".tsv": "text/tab-separated-values",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".doc": "application/msword",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".tiff": "image/tiff",
    ".zip": "application/zip",
}


def sanitize_filename(filename: str) -> str:
    """
    Sanitize filename by removing dangerous path separators and special characters.
    """
    clean = os.path.basename(filename)
    clean = re.sub(r"[^\w\s\.-]", "_", clean).strip()
    clean = re.sub(r"\s+", "_", clean)
    if not clean or clean.startswith("."):
        clean = f"document_{clean.lstrip('.')}" if clean.lstrip(".") else "document"
    return clean


def get_file_extension(filename: str) -> str:
    """
    Extract lowercase extension including the dot.
    """
    return Path(filename).suffix.lower()


def validate_file_extension(filename: str) -> str:
    """
    Validate that the file extension is allowed.
    Raises HTTPException(415) if not supported.
    """
    ext = get_file_extension(filename)
    if not ext or ext not in ALLOWED_EXTENSIONS:
        allowed_list = ", ".join(sorted(ALLOWED_EXTENSIONS))
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail=f"Unsupported file type '{ext}'. Allowed types: {allowed_list}",
        )
    return ext


def detect_document_type(filename: str) -> DocumentType:
    """
    Auto-detect document type from filename keywords.
    """
    lower = filename.lower()
    if any(k in lower for k in ["doa", "received", "packing list", "packing_list", "delivery note", "delivery_note", "delivery", "asset register", "asset_register"]):
        return DocumentType.DELIVERY
    if "hostname" in lower and any(k in lower for k in ["serial", "location", "device"]):
        return DocumentType.DELIVERY
    if any(k in lower for k in ["bom", "bill_of_materials", "bill-of-materials", "materials"]):
        return DocumentType.BOM
    if any(k in lower for k in ["p2p", "point_to_point", "point-to-point", "port_to_port", "interconnect"]):
        return DocumentType.P2P
    if any(k in lower for k in ["rack", "elevation", "layout", "cabinet"]):
        return DocumentType.RACK_LAYOUT
    if any(k in lower for k in ["power", "pdu", "ups", "electric", "feed"]):
        return DocumentType.POWER
    if any(k in lower for k in ["cable", "cabling", "patch", "wiring", "fiber"]):
        return DocumentType.CABLE
    return DocumentType.OTHER


def detect_document_type_from_content(file_path: Union[Path, str], filename: Optional[str] = None) -> DocumentType:
    """
    Auto-detect document type inspecting filename, sheet names, and column headers.
    Special detection for 'delivery':
    - Sheet or column names containing: 'DoA', 'Received', 'Packing list', or 'Delivery note'
    - Or columns with Serial No. + Rack + RU
    """
    path = Path(file_path)
    fn = filename or path.name
    type_from_name = detect_document_type(fn)
    if type_from_name == DocumentType.DELIVERY:
        return DocumentType.DELIVERY

    ext = path.suffix.lower()

    def check_headers_for_delivery(headers_list) -> bool:
        norm_headers = [re.sub(r"[_\s\-]+", " ", str(h).strip().lower()) for h in headers_list if h is not None]
        # 1. 'DoA', 'Received', 'Packing list', 'Delivery note'
        delivery_keywords = ["doa", "received", "packing list", "delivery note"]
        for h in norm_headers:
            if any(k in h for k in delivery_keywords):
                return True

        # 2. Serial No. + Rack + RU columns
        has_serial = any(any(s in h for s in ["serial", "serial no", "serial number", "s/n", "sn", "serial #"]) for h in norm_headers)
        has_rack = any(any(r in h for r in ["rack", "rack id", "cabinet", "cabinet id"]) for h in norm_headers)
        has_ru = any(re.search(r"\b(u|ru|u position|rack unit|slot)\b", h) for h in norm_headers)
        if has_serial and has_rack and has_ru:
            return True
        return False

    try:
        if ext in [".xlsx", ".xls", ".xlsm"]:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            # Check sheet names
            delivery_keywords = ["doa", "received", "packing list", "delivery note", "delivery", "asset register"]
            for sname in wb.sheetnames:
                s_lower = sname.lower()
                if any(k in s_lower for k in delivery_keywords):
                    wb.close()
                    return DocumentType.DELIVERY

            # Check header rows across sheets
            for sname in wb.sheetnames:
                sheet = wb[sname]
                for row in sheet.iter_rows(max_row=10, values_only=True):
                    non_empty = [c for c in row if c is not None and str(c).strip()]
                    if len(non_empty) >= 2:
                        if check_headers_for_delivery(non_empty):
                            wb.close()
                            return DocumentType.DELIVERY
            wb.close()
        elif ext in [".csv", ".tsv"]:
            import csv
            sep = "\t" if ext == ".tsv" else ","
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f, delimiter=sep)
                for _ in range(10):
                    row = next(reader, None)
                    if row is None:
                        break
                    non_empty = [c for c in row if c.strip()]
                    if len(non_empty) >= 2:
                        if check_headers_for_delivery(non_empty):
                            return DocumentType.DELIVERY
    except Exception:
        pass

    return type_from_name


def get_mime_type_for_file(file_path: Union[Path, str]) -> str:
    """
    Get appropriate MIME type for download streaming.
    """
    ext = Path(file_path).suffix.lower()
    return EXTENSION_MIME_MAP.get(ext, "application/octet-stream")


async def save_upload_file(
    upload_file: UploadFile,
    dest_path: Path,
    max_size_mb: Optional[int] = None,
) -> Tuple[int, str]:
    """
    Stream upload file to disk while calculating SHA-256.
    If max_size_mb is provided, enforces max size; otherwise allows unlimited size.
    Returns (file_size_bytes, sha256_checksum).
    """
    dest_path.parent.mkdir(parents=True, exist_ok=True)
    max_bytes = max_size_mb * 1024 * 1024 if max_size_mb else None
    total_bytes = 0
    hasher = hashlib.sha256()

    try:
        with open(dest_path, "wb") as f:
            chunk_size = 1024 * 1024  # 1MB chunks
            while True:
                chunk = await upload_file.read(chunk_size)
                if not chunk:
                    break
                total_bytes += len(chunk)
                if max_bytes and total_bytes > max_bytes:
                    raise HTTPException(
                        status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                        detail=f"File exceeds maximum allowed size of {max_size_mb} MB",
                    )
                hasher.update(chunk)
                f.write(chunk)
    except Exception:
        if dest_path.exists():
            dest_path.unlink(missing_ok=True)
        raise

    return total_bytes, hasher.hexdigest()


def delete_file_safely(file_path: Path) -> bool:
    """
    Safely delete a file from disk if it exists.
    """
    try:
        if file_path.exists():
            file_path.unlink()
            return True
        return False
    except OSError:
        return False
